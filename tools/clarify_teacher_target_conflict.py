from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pathlib import Path

SRC=Path('docs/research/conflict_client_sweep_20261020/presentation/Federated_LEM2_Teacher_Questions_Complete_20261023.pptx')
OUT=SRC.with_name('Federated_LEM2_Teacher_Questions_Complete_20261023.clarified.tmp.pptx')
prs=Presentation(SRC)
FONT='Noto Sans CJK TC'; MATH='Cambria Math'
C={'dark':'1F2933','teal':'2B7A78','tealL':'E7F2F1','blue':'4A78A8','blueL':'EAF2F8','orange':'C96A3D','orangeL':'F8EAE3','green':'4C8A61','greenL':'E8F3E8','red':'B65252','redL':'FBEAEA','gray':'6B7780','grayL':'F5F7F8','line':'D9E2E7','gold':'A98025','goldL':'F6F0DC','white':'FFFFFF','purple':'765A9D','purpleL':'F0ECF6'}
def rgb(h): return RGBColor.from_string(h)
def clear_slide(s):
    for sh in list(s.shapes):
        s.shapes._spTree.remove(sh._element)
    s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(C['white'])
def tx(s,text,x,y,w,h,size=14,color='1F2933',bold=False,align=PP_ALIGN.LEFT,valign=MSO_ANCHOR.TOP,font=FONT):
    box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0; tf.vertical_anchor=valign
    p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=str(text); r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=rgb(color); return box
def rect(s,x,y,w,h,fill,linec=None,round=True):
    shp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); shp.fill.solid(); shp.fill.fore_color.rgb=rgb(fill); shp.line.color.rgb=rgb(linec or fill); return shp
def line(s,x1,y1,x2,y2,color='D9E2E7',width=1):
    ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); ln.line.color.rgb=rgb(color); ln.line.width=Pt(width); return ln
def title(s,t,sub=None):
    tx(s,t,0.62,0.34,12.05,0.48,24,C['dark'],True)
    if sub: tx(s,sub,0.64,0.86,12.0,0.34,12.5,C['teal'])
def footer(s,n):
    line(s,0.55,7.07,12.75,7.07,C['line'],0.7); tx(s,'Federated LEM2 · Target / conflict intuition · 2026-10-24',0.62,7.13,10.5,0.18,8,C['gray']); tx(s,str(n),12.2,7.13,0.4,0.18,8,C['gray'],False,PP_ALIGN.RIGHT)
def card(s,x,y,w,h,head,body,fill='F5F7F8',accent='2B7A78',hs=14,bs=11.5):
    rect(s,x,y,w,h,fill); rect(s,x+0.12,y+0.15,0.055,h-0.3,accent,accent,False); tx(s,head,x+0.32,y+0.15,w-0.48,0.30,hs,C['dark'],True); tx(s,body,x+0.32,y+0.56,w-0.50,h-0.66,bs,C['dark'])
def takeaway(s,text):
    rect(s,0.72,6.31,11.9,0.56,C['tealL']); tx(s,'一句話：',0.92,6.46,0.85,0.2,10.5,C['teal'],True); tx(s,text,1.75,6.42,10.35,0.26,10.7,C['dark'])
def formula(s,text,x,y,w,h,size=18,fill='F5F7F8',color='1F2933'):
    rect(s,x,y,w,h,fill,C['line']); tx(s,text,x+0.12,y+0.11,w-0.24,h-0.22,size,color,True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE,MATH)
def circle(s,x,y,d,fill,text='',tcolor='1F2933',fs=9,bold=True):
    shp=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d)); shp.fill.solid(); shp.fill.fore_color.rgb=rgb(fill); shp.line.color.rgb=rgb(fill)
    if text: tx(s,text,x,y+d*0.25,d,d*0.38,fs,tcolor,bold,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
    return shp

# 16 — what targets are
s=prs.slides[15]; clear_slide(s); title(s,'Target 1–5 到底是什麼？','不是 5 個 factors；是每個 dataset 從 20% holdout 挑 5 筆「測試資料」做 X-ray 回查')
rect(s,0.78,1.42,11.8,1.0,C['blueL']); tx(s,'一個 Target = 一筆完整 test row x',1.0,1.68,4.1,0.3,18,C['blue'],True); tx(s,'它帶著所有 factor values + true decision',5.35,1.70,6.9,0.3,14,C['dark'],True)
for i,(lab,sub,fill,ac) in enumerate([('Target 1','代表 corrected',C['greenL'],C['green']),('Target 2','另一種 corrected',C['tealL'],C['teal']),('Target 3','代表 harmed',C['redL'],C['red']),('Target 4','另一種 harmed',C['orangeL'],C['orange']),('Target 5','缺席 / unchanged',C['goldL'],C['gold'])]):
    x=0.78+i*2.36; rect(s,x,2.82,2.08,1.35,fill); tx(s,lab,x+0.08,3.05,1.92,0.28,14,ac,True,PP_ALIGN.CENTER); tx(s,sub,x+0.12,3.52,1.84,0.32,10.4,C['dark'],False,PP_ALIGN.CENTER)
card(s,0.92,4.57,5.48,1.25,'為什麼只挑 5 筆？','先用少量代表 case 把「為什麼翻轉」講清楚；它不是顯著性 sample。',C['grayL'],C['gray'],13,11.2)
card(s,6.82,4.57,5.48,1.25,'後面怎麼擴大？','同一套回查流程再掃所有 conflict cases，才得到 908、45、43 這些全體統計。',C['tealL'],C['teal'],13,11.2)
takeaway(s,'Target 1–5 是五筆 test objects 的解剖案例，不是五個變數，也不是五種模型。'); footer(s,16)

# 17 — x-ray pipeline
s=prs.slides[16]; clear_slide(s); title(s,'一筆 Target 怎麼看？像做一張 prediction X-ray','先看資料長什麼樣，再看它撞到哪些 rules，最後才看各 policy 怎麼選 winner')
steps=[('① 輸入 x','factors\n+ true class',C['blueL'],C['blue']),('② Matching Rules','x 滿足哪些\nglobal rules',C['tealL'],C['teal']),('③ Candidate','這些 rules 提出\n哪些 decisions',C['goldL'],C['gold']),('④ Evidence','每個 decision 的\nT / U / P / R',C['orangeL'],C['orange']),('⑤ Winner','STD / PS / GS / PG\n最後選誰',C['greenL'],C['green'])]
for i,(h,b,f,a) in enumerate(steps):
    x=0.48+i*2.55; rect(s,x,1.88,2.12,1.62,f); tx(s,h,x+0.08,2.13,1.96,0.32,14,a,True,PP_ALIGN.CENTER); tx(s,b,x+0.12,2.68,1.88,0.48,10.5,C['dark'],False,PP_ALIGN.CENTER)
    if i<4: line(s,x+2.14,2.70,x+2.48,2.70,C['line'],1.5)
formula(s,'M(x)={r : x⊨Ante(r)}',0.78,3.95,3.6,0.65,15,C['grayL'])
formula(s,'Candidate(x)={d_r : r∈M(x)}',4.86,3.95,3.6,0.65,14,C['grayL'])
formula(s,'T_d, U_d, P_d=U_d/T_d, R_d=1-P_d',8.94,3.95,3.6,0.65,13,C['grayL'])
card(s,1.3,4.92,10.75,1.02,'先問這題','True class 有沒有進 Candidate(x)？YES → 才是 ranking / conflict-resolution；NO → 任何 voting 都救不了。',C['redL'],C['red'],12.5,11.2)
takeaway(s,'回查不是只看 prediction；是把「rule → candidate → evidence → score → winner」每一層打開。'); footer(s,17)

# 18 — Balance visual
s=prs.slides[17]; clear_slide(s); title(s,'Balance Target 1：Standard 為什麼錯，去重後為什麼會修正？','同一 training object 被多條 L-rules 重複計價，讓 L 的 T 值被放大')
card(s,0.62,1.30,2.62,2.05,'Target x','rd=5 · rw=1\nld=3 · lw=1\n\nTrue = R',C['blueL'],C['blue'],14,12.2)
card(s,3.48,1.30,3.55,2.05,'Matching rules','r1 → R : support 18\nr2 → L : support 3\nr3 → L : support 16\n\nCandidate={L,R}',C['tealL'],C['teal'],14,11.3)
rect(s,7.30,1.30,5.35,2.05,C['grayL']); tx(s,'L：19 次「被計分」≠ 19 個不同 objects',7.55,1.50,4.85,0.30,13,C['red'],True,PP_ALIGN.CENTER)
for i in range(16):
    row=i//8; col=i%8; circle(s,7.60+col*0.47,1.95+row*0.46,0.27,C['blueL'],str(i+1),'4A78A8',6.5,False)
for j in range(3): circle(s,11.52+j*0.31,2.19,0.24,C['redL'],'↺',C['red'],7.5,True)
tx(s,'16 unique objects + 3 duplicate occurrences',7.66,2.92,4.65,0.22,10,C['dark'],True,PP_ALIGN.CENTER)
formula(s,'L: T_L=19, U_L=16',0.82,3.72,3.7,0.64,16,C['redL'],C['red'])
formula(s,'R: T_R=18, U_R=18',4.82,3.72,3.7,0.64,16,C['greenL'],C['green'])
formula(s,'P_L=16/19=.842 ; P_R=1',8.82,3.72,3.7,0.64,14,C['grayL'])
card(s,0.88,4.68,5.55,1.02,'Standard 比 T','19 > 18 → L 勝 → True=R → WRONG',C['redL'],C['red'],13.5,11.7)
card(s,6.88,4.68,5.55,1.02,'Population Standard 比 U','16 < 18 → R 勝 → CORRECTED',C['greenL'],C['green'],13.5,11.7)
takeaway(s,'PS 不是給 L 負分；只是把同一 training object 的重複 contribution 合併成一次。'); footer(s,18)

# 19 — how to read target cards
s=prs.slides[18]; clear_slide(s); title(s,'接下來 Target 1–5 三頁怎麼讀？','每一列只回答三件事：真實答案、Population 有沒有改 winner、true class 有沒有被 rule system 提出')
card(s,0.78,1.42,3.75,3.42,'A · True + Factors','True = 這筆 holdout 的真實 decision。\n\nFactors = 這筆資料本身的條件值。\n\nC = 當下把 80% training split 成幾個 clients。',C['blueL'],C['blue'],14,12)
card(s,4.80,1.42,3.75,3.42,'B · STD→PS / GS→PG','Corrected：wrong → true\nHarmed：true → wrong\nChanged / same：winner 變了，但 correctness 沒變\nUnchanged：winner 沒變',C['tealL'],C['teal'],14,11.6)
card(s,8.82,1.42,3.75,3.42,'C · True cand','YES：true class 至少有一條 matching rule，可以靠 scoring 去競爭。\n\nNO：true class 根本沒進 Candidate(x)，任何 voting policy 都不可能救。',C['redL'],C['red'],14,11.6)
rect(s,1.30,5.14,10.7,0.72,C['goldL']); tx(s,'閱讀順序：True cand → old/new winner → 再回頭看 T/U/P/R；不要只看「Corrected / Harmed」標籤。',1.52,5.34,10.25,0.27,12,C['gold'],True,PP_ALIGN.CENTER)
takeaway(s,'Target cards 是 case-level mechanism evidence；不是在說某一個 Target 比另一個 Target「比較好」。'); footer(s,19)

# 20 — 15 cases role
s=prs.slides[19]; clear_slide(s); title(s,'15 個 Target 的角色：不是做統計，而是把四種情況「看懂」','3 datasets × 5 cases，刻意涵蓋 corrected、harmed、unchanged、true-class missing')
for i,(h,b,f,a) in enumerate([
('Corrected','baseline 錯 → population 修正',C['greenL'],C['green']),('Harmed','baseline 對 → population 修壞',C['redL'],C['red']),('Changed / same','winner 變，但 correctness 不變',C['goldL'],C['gold']),('Unchanged','penalty 有算，但不足以翻轉',C['blueL'],C['blue']),('True missing','true class 根本沒進 Candidate',C['orangeL'],C['orange'])]):
    x=0.58+i*2.48; rect(s,x,1.75,2.15,1.75,f); tx(s,h,x+0.08,2.06,1.99,0.28,13,a,True,PP_ALIGN.CENTER); tx(s,b,x+0.12,2.58,1.91,0.48,10.3,C['dark'],False,PP_ALIGN.CENTER)
formula(s,'15 selected cases  →  理解機制',1.05,4.15,4.8,0.66,17,C['grayL'])
formula(s,'全部 conflict cases  →  統計機制出現多少次',7.0,4.15,5.0,0.66,16,C['grayL'])
card(s,1.25,5.12,10.82,0.82,'下一步','所以先看 21–23 頁三個 dataset 的 5 個 target；看懂後，再把同一套判讀規則套到所有 conflict occurrences。',C['tealL'],C['teal'],12.5,11)
takeaway(s,'15 cases 是「顯微鏡」；908 conflict scan 才是「把整張樣本掃完」。'); footer(s,20)

# 24 — conflict not six methods
s=prs.slides[23]; clear_slide(s); title(s,'先回答：908 跟「六種方法」完全不是同一件事','Conflict 在 prediction policy 之前就決定；六種方法只是同一個 conflict 發生後，用不同分數選 winner')
card(s,0.70,1.38,5.95,3.65,'先建立 rule system · 一次','Dataset + C\n→ 80% 分 clients\n→ Local LEM2\n→ UNION global rules\n→ test x 找 M(x)\n→ Candidate(x)\n→ |Candidate(x)|>1 ? = Conflict',C['blueL'],C['blue'],15,11.8)
rect(s,6.98,1.38,5.62,3.65,C['grayL']); tx(s,'Conflict 已經存在後，才跑 6 種 policy',7.22,1.67,5.14,0.30,14,C['dark'],True,PP_ALIGN.CENTER)
for i,(lab,fill,ac) in enumerate([('FIRST',C['orangeL'],C['orange']),('SIMPLE',C['goldL'],C['gold']),('STD',C['blueL'],C['blue']),('GS',C['greenL'],C['green']),('PS',C['tealL'],C['teal']),('PG',C['redL'],C['red'])]):
    col=i%3; row=i//3; rect(s,7.28+col*1.60,2.30+row*0.88,1.34,0.62,fill); tx(s,lab,7.28+col*1.60,2.49+row*0.88,1.34,0.20,10.5,ac,True,PP_ALIGN.CENTER)
formula(s,'Conflict(x) ⇔ |Candidate(x)| > 1',1.55,5.30,4.55,0.62,16,C['grayL'])
formula(s,'這個公式沒有 policy p',7.20,5.30,4.20,0.62,17,C['redL'],C['red'])
takeaway(s,'六種方法不會把 conflict 從 1 變成 6；它們共用同一個 conflict set。'); footer(s,24)

# 25 — exact 908 and why large
s=prs.slides[24]; clear_slide(s); title(s,'908 怎麼來？為什麼看起來這麼大？','它是 12 個 structural configurations 的累積；同一 holdout row 在不同 C 下可以被重複計為 conflict occurrence')
xs=[2.65,4.38,6.11,7.84,9.57]; headers=['C=2','C=3','C=4','C=5','小計']
for x,h in zip(xs,headers): rect(s,x,1.42,1.43,0.52,C['tealL']); tx(s,h,x,1.58,1.43,0.18,10.8,C['teal'],True,PP_ALIGN.CENTER)
rows=[('Balance','47\n37.3%','62\n49.2%','74\n58.7%','83\n65.9%','266',C['blueL'],C['blue']),('Car','85\n24.6%','134\n38.7%','169\n48.8%','210\n60.7%','598',C['greenL'],C['green']),('Hayes','3\n11.5%','14\n53.8%','10\n38.5%','17\n65.4%','44',C['orangeL'],C['orange'])]
for r,(name,*vals,fill,ac) in enumerate(rows):
    y=2.12+r*0.95; rect(s,0.80,y,1.55,0.72,fill); tx(s,name,0.80,y+0.23,1.55,0.20,11.5,ac,True,PP_ALIGN.CENTER)
    for j,val in enumerate(vals): rect(s,xs[j],y,1.43,0.72,'FAFBFC','D9E2E7'); tx(s,val,xs[j],y+0.10,1.43,0.42,10.5,C['dark'],True,PP_ALIGN.CENTER)
formula(s,'266 + 598 + 44 = 908 conflict occurrences',1.02,5.05,5.38,0.62,16,C['goldL'],C['gold'])
card(s,6.72,4.88,5.65,1.02,'為什麼不是 908 個 unique rows？','每個 dataset 的同一 20% holdout 會在 C=2,3,4,5 各跑一次；同一 row 可在多個 C 都 conflict。',C['redL'],C['red'],12.5,10.6)
rect(s,1.05,5.89,11.3,0.34,C['grayL']); tx(s,'總共 configuration-level test evaluations = 126×4 + 346×4 + 26×4 = 1,992；其中 908 次 conflict ≈ 45.6%。',1.20,5.97,11.0,0.18,10.5,C['dark'],True,PP_ALIGN.CENTER)
takeaway(s,'908 是 3 datasets × 4 個 C 的累積 workload；不是 ×6 policies。'); footer(s,25)

# 26 — why C increases conflict
s=prs.slides[25]; clear_slide(s); title(s,'為什麼 C 從 2 增加到 5，Conflict 通常會上升？','目前資料最一致的結構鏈：client fragmentation → rules 更一般 → matching overlap 增加')
chain=[('C ↑','client 數增加',C['blueL'],C['blue']),('|U_c| ↓','每 client\n資料變少',C['tealL'],C['teal']),('Rule shorter','local rule\n較短/較一般',C['orangeL'],C['orange']),('Matches ↑','每筆 x\nmatch 更多 rules',C['goldL'],C['gold']),('Decision overlap ↑','更多不同\ndecisions 同時出現',C['redL'],C['red']),('Conflict ↑','|Candidate|>1\n更常發生',C['greenL'],C['green'])]
for i,(a,b,f,ac) in enumerate(chain):
    x=0.42+i*2.10; rect(s,x,1.75,1.75,1.55,f); tx(s,a,x+0.06,2.02,1.63,0.30,14,ac,True,PP_ALIGN.CENTER); tx(s,b,x+0.10,2.52,1.55,0.46,10.3,C['dark'],False,PP_ALIGN.CENTER)
    if i<5: line(s,x+1.77,2.52,x+2.04,2.52,C['line'],1.4)
card(s,0.82,3.77,5.70,1.72,'直覺','同樣 80% training，切成更多 clients 後，每個 client 只看較少 rows。Local LEM2 更容易學到「少幾個條件也成立」的較寬 rules；UNION 後這些 rules 更容易互相重疊。',C['grayL'],C['gray'],13.3,11)
card(s,6.82,3.77,5.70,1.72,'但不是定理','Hayes test 只有 26 筆，因此 C=3→4 conflict rate 會從 53.8% 回落到 38.5%；目前只能說整體與 endpoint 趨勢支持這條機制。',C['redL'],C['red'],13.3,11)
takeaway(s,'Conflict 上升是 rule structure / overlap 改變，不是「方法變成 6 個」。'); footer(s,26)

# 27 — actual evidence
s=prs.slides[26]; clear_slide(s); title(s,'數據真的有沿著這條鏈走嗎？Car / Balance：有','C=2 → C=5 時，rule 變短、matching rules / test 增加、conflict rate 同步上升')
cards=[('Balance Scale','約 250 → 100','2.98 → 2.47','2.43 → 3.74','37.3% → 65.9%',C['blueL'],C['blue']),('Car Evaluation','691 → 276','4.40 → 3.66','1.74 → 3.16','24.6% → 60.7%',C['greenL'],C['green'])]
for i,(name,rows,cond,match,cr,fill,ac) in enumerate(cards):
    x=0.72+i*6.05; rect(s,x,1.38,5.55,4.02,fill); tx(s,name,x+0.25,1.66,5.05,0.32,17,ac,True,PP_ALIGN.CENTER)
    metrics=[('平均 train rows / client',rows),('Local rule 平均條件數',cond),('Matching rules / test',match),('Conflict rate',cr)]
    for k,(lab,val) in enumerate(metrics):
        y=2.30+k*0.72; tx(s,lab,x+0.35,y,3.0,0.25,11,C['dark']); tx(s,val,x+3.25,y,1.80,0.25,13,ac,True,PP_ALIGN.RIGHT)
formula(s,'C ↑  →  local rule conditions ↓  →  matching rules / test ↑  →  conflict rate ↑',1.25,5.62,10.82,0.62,15,C['grayL'])
takeaway(s,'這就是我們目前說「fragmentation 會增加 conflict workload」最直接的結構證據。'); footer(s,27)

# 28 — exclude simplistic alternatives
s=prs.slides[27]; clear_slide(s); title(s,'不是「規則越多所以 conflict 越多」；也不一定是 reduct 不同','兩個反例幫我們把原因收斂到 rule generality / overlap')
card(s,0.78,1.45,5.75,3.72,'反例 1 · Balance：global rules 反而變少','Global rule count：C=2 → 209；C=5 → 183。\n\n但 conflict rate：37.3% → 65.9%。\n\n所以 rule 總數增加不是必要條件；更關鍵的是每條 rule 覆蓋多寬、彼此重疊多少。',C['blueL'],C['blue'],15,12)
card(s,6.82,1.45,5.75,3.72,'反例 2 · Car / Balance：reduct instability = 0','Car、Balance 在 C=2..5 的 reduct instability 都是 0，\n但 conflict 仍然大幅增加。\n\n因此也不能把 conflict 全歸因於「不同 clients 選到不同 reduct」。',C['orangeL'],C['orange'],15,12)
rect(s,1.18,5.48,11.0,0.58,C['tealL']); tx(s,'目前最合理：fragmentation → rule generalization → matching overlap → cross-decision conflict',1.42,5.66,10.5,0.20,12.5,C['teal'],True,PP_ALIGN.CENTER)
takeaway(s,'真正要追的是「rule 結構怎麼變」，不是只數 rule 有幾條。'); footer(s,28)

# 29 — what policies affect
s=prs.slides[28]; clear_slide(s); title(s,'六種方法到底改了什麼？它們只改「Conflict 發生後的排序」','所以 908 是 conflict workload；43 才是 Population correction 真正翻動 winner 的次數')
formula(s,'908 = structural conflict occurrences',0.85,1.30,5.6,0.66,18,C['blueL'],C['blue'])
formula(s,'43 = population-induced winner changes',6.85,1.30,5.6,0.66,18,C['tealL'],C['teal'])
card(s,0.80,2.28,3.75,2.75,'908 裡的 45','true class 不在 Candidate(x)\n→ voting 根本救不了\n→ Coverage / Induction failure',C['redL'],C['red'],14,11.8)
card(s,4.80,2.28,3.75,2.75,'STD → PS','Winner changed 16 次\n6 corrected\n9 harmed\n1 wrong→wrong',C['blueL'],C['blue'],14,11.8)
card(s,8.80,2.28,3.75,2.75,'GS → PG','Winner changed 27 次\n10 corrected\n13 harmed\n4 same correctness',C['greenL'],C['green'],14,11.8)
formula(s,'43 / 908 = 4.74%：Population correction 只翻動少數 borderline conflicts',1.45,5.38,10.45,0.66,15.5,C['goldL'],C['gold'])
takeaway(s,'先有 conflict，再有 resolver；六方法不改 conflict 數，只改同一 conflict 的 winner。'); footer(s,29)

# 30 — asymmetry and final chain
s=prs.slides[29]; clear_slide(s); title(s,'最後才進 Redundancy Asymmetry：解釋「為什麼有時救對、有時修壞」','它不是第七種 prediction method，而是 conflict evidence structure 的 diagnostic')
formula(s,'R_d = 1 - U_d/T_d',0.90,1.28,3.65,0.65,18,C['grayL'])
formula(s,'A(x)=R_competitor - R_true',4.85,1.28,3.65,0.65,17,C['grayL'])
formula(s,'43/43 changed winner → lower redundancy',8.80,1.28,3.65,0.65,14,C['tealL'],C['teal'])
card(s,0.78,2.24,5.72,2.65,'A(x) > 0 · 較像 Corrected','錯誤 competitor 比 true class 更重複。\n例 Car T2：R_acc=.300，R_good=0，A=+.300。\nPenalty 壓低 acc → good 反超。',C['greenL'],C['green'],14,11.4)
card(s,6.82,2.24,5.72,2.65,'A(x) < 0 · 較像 Harmed','true class 自己反而更重複。\n例 Car T4：R_unacc=0，R_acc=.4179，A=-.4179。\nPenalty 壓掉正確 acc → unacc 反超。',C['redL'],C['red'],14,11.4)
rect(s,0.95,5.25,11.45,0.62,C['purpleL']); tx(s,'C → Fragmentation → Rule generality / overlap → Conflict → Redundancy structure → Resolver → Performance',1.18,5.45,11.0,0.22,12.3,C['purple'],True,PP_ALIGN.CENTER)
takeaway(s,'研究重點不是「conflict 多不多」而已，而是 conflict 裡哪個 decision 的 support 被重複放大。'); footer(s,30)

for idx,s in enumerate(prs.slides,1):
    for sh in s.shapes:
        if hasattr(sh,'text'):
            t=sh.text.strip()
            if t.isdigit() and sh.top > Inches(6.8) and sh.left > Inches(11.5):
                sh.text=str(idx)

prs.save(OUT)
OUT.replace(SRC)
print(SRC)
print('slides',len(prs.slides),'tables',sum(1 for s in prs.slides for sh in s.shapes if getattr(sh,'has_table',False)))
