from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pathlib import Path

SRC=Path('docs/research/conflict_client_sweep_20261020/presentation/Federated_LEM2_Teacher_Questions_Complete_20261023.pptx')
OUT=SRC.with_name('Federated_LEM2_Teacher_Questions_Complete_20261023.enriched.tmp.pptx')
prs=Presentation(SRC)
blank=prs.slide_layouts[6]
FONT='Noto Sans CJK TC'
MATH='Cambria Math'
C={'dark':'1F2933','teal':'2B7A78','tealL':'E7F2F1','blue':'4A78A8','blueL':'EAF2F8','orange':'C96A3D','orangeL':'F8EAE3','green':'4C8A61','greenL':'E8F3E8','red':'B65252','redL':'FBEAEA','gray':'6B7780','grayL':'F5F7F8','line':'D9E2E7','gold':'A98025','goldL':'F6F0DC','white':'FFFFFF'}
def rgb(h): return RGBColor.from_string(h)
def tx(s,text,x,y,w,h,size=14,color='1F2933',bold=False,align=PP_ALIGN.LEFT,valign=MSO_ANCHOR.TOP,font=FONT):
    box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0; tf.vertical_anchor=valign
    p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=str(text); r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=rgb(color); return box
def rect(s,x,y,w,h,fill,line=None,round=True):
    shp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); shp.fill.solid(); shp.fill.fore_color.rgb=rgb(fill); shp.line.color.rgb=rgb(line or fill); return shp
def line(s,x1,y1,x2,y2,color='D9E2E7',width=1):
    ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); ln.line.color.rgb=rgb(color); ln.line.width=Pt(width); return ln
def title(s,t,sub=None):
    tx(s,t,0.62,0.34,12.05,0.48,24,C['dark'],True)
    if sub: tx(s,sub,0.64,0.86,12.0,0.34,12.5,C['teal'])
def footer(s,n=0):
    line(s,0.55,7.07,12.75,7.07,C['line'],0.7); tx(s,'Federated LEM2 · Target backtrace explanation · 2026-10-24',0.62,7.13,10.5,0.18,8,C['gray']); tx(s,str(n),12.2,7.13,0.4,0.18,8,C['gray'],False,PP_ALIGN.RIGHT)
def card(s,x,y,w,h,head,body,fill='F5F7F8',accent='2B7A78',hs=14,bs=11.5):
    rect(s,x,y,w,h,fill); rect(s,x+0.12,y+0.15,0.055,h-0.3,accent,accent,False); tx(s,head,x+0.32,y+0.15,w-0.48,0.30,hs,C['dark'],True); tx(s,body,x+0.32,y+0.56,w-0.50,h-0.66,bs,C['dark'])
def takeaway(s,text):
    rect(s,0.72,6.31,11.9,0.56,C['tealL']); tx(s,'結論：',0.92,6.46,0.8,0.2,10.5,C['teal'],True); tx(s,text,1.65,6.42,10.55,0.26,10.7,C['dark'])
def formula(s, text, x,y,w,h, size=18, fill='F5F7F8', color='1F2933'):
    rect(s,x,y,w,h,fill,C['line']); tx(s,text,x+0.15,y+0.13,w-0.3,h-0.25,size,color,True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE,MATH)
def new_slide():
    s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(C['white']); return s

new=[]
s=new_slide(); new.append(s)
title(s,'Target 回查：一筆測試資料到底怎麼一路算到最後？','先把「回查」定義成可重現的計算鏈，而不是只看最後 prediction')
steps=[('① Target x','20% holdout 的一筆資料',C['blueL'],C['blue']),('② Matching rules','M(x) = 符合 antecedent 的 global rules',C['tealL'],C['teal']),('③ Candidate decisions','Candidate(x) = {d_r : r∈M(x)}',C['goldL'],C['gold']),('④ Evidence','每個 decision 算 T_d、U_d、P_d、R_d',C['orangeL'],C['orange']),('⑤ Policy scores','STD / PS / GS / PG 分別計分',C['greenL'],C['green']),('⑥ Outcome','old winner → new winner → Corrected / Harmed',C['redL'],C['red'])]
for i,(h,b,f,a) in enumerate(steps):
    col=i%3; row=i//3; card(s,0.7+col*4.18,1.45+row*2.05,3.75,1.55,h,b,f,a,13.5,11)
    if i<2: line(s,4.45+col*4.18,2.22,4.78+col*4.18,2.22,C['line'],1.4)
    if 3<=i<5: line(s,4.45+(i-3)*4.18,4.27,4.78+(i-3)*4.18,4.27,C['line'],1.4)
formula(s,'若 y ∉ Candidate(x) → 這不是 voting 問題，而是 Rule Coverage / Induction Failure',1.35,5.48,10.65,0.62,15,C['redL'],C['red'])
takeaway(s,'後面的 15 targets、908 conflicts、43 winner changes，全部都是沿著這條鏈計算。'); footer(s)

s=new_slide(); new.append(s)
title(s,'Step 1 — 先找 M(x) 與 Candidate(x)','Balance Scale · Target 1：先確認「哪些 rules 有資格參與這筆 prediction」')
card(s,0.72,1.35,3.05,2.75,'Target x','right_distance = 5\nright_weight = 1\nleft_distance = 3\nleft_weight = 1\n\nTrue decision = R',C['blueL'],C['blue'],14,12)
card(s,4.05,1.35,4.25,2.75,'Matching rules M(x)','r1: right_distance=5 AND left_weight=1 → R\n\nr2: right_weight=1 AND right_distance=5 AND left_distance=3 → L\n\nr3: right_weight=1 AND left_distance=3 → L',C['tealL'],C['teal'],14,10.6)
card(s,8.58,1.35,4.02,2.75,'Candidate decisions','M_R(x) = {r1}\nM_L(x) = {r2, r3}\n\nCandidate(x) = {L, R}\n\n因為 |Candidate(x)| = 2，所以這筆是真正的 prediction-time conflict。',C['goldL'],C['gold'],14,11.2)
formula(s,'M(x) = { r ∈ R_G : x ⊨ Ante(r) }',1.05,4.52,5.35,0.7,18,C['grayL'])
formula(s,'Candidate(x) = { d_r : r ∈ M(x) }',6.9,4.52,5.35,0.7,18,C['grayL'])
takeaway(s,'回查的第一件事不是算分數，而是先知道「有哪些 decision 真的被 matching rules 提出」。'); footer(s)

s=new_slide(); new.append(s)
title(s,'Step 2 — 每條 Rule 的 support 從哪裡來？','只回查 80% training objects；20% holdout 永遠不參與 support / weight')
formula(s,'S_G(r) = { z ∈ U_train : z ⊨ Ante(r) ∧ D(z)=d_r }',0.9,1.32,11.55,0.72,17,C['tealL'],C['teal'])
card(s,0.85,2.32,3.65,2.05,'R 的 evidence','r1 → R\n|S_G(r1)| = 18\n\n因此：\nT_R = 18\nU_R = 18',C['greenL'],C['green'],14,12)
card(s,4.85,2.32,3.65,2.05,'L 的 evidence','r2 → L : support = 3\nr3 → L : support = 16\n\nStandard 先直接加總：\nT_L = 3 + 16 = 19',C['orangeL'],C['orange'],14,11.6)
card(s,8.85,2.32,3.65,2.05,'重複在哪裡？','r2 的 3 個 support objects 全部也出現在 r3。\n\n因此 19 次 support occurrences 裡，真正不同的 objects 只有 16 個：\nU_L = 16',C['redL'],C['red'],14,10.8)
formula(s,'T_d = Σ_{r∈M_d(x)} |S_G(r)|',0.9,4.72,3.6,0.72,16,C['grayL'])
formula(s,'U_d = | ⋃_{r∈M_d(x)} S_G(r) |',4.85,4.72,3.6,0.72,16,C['grayL'])
formula(s,'P_d = U_d / T_d   ;   R_d = 1 - P_d',8.8,4.72,3.6,0.72,15,C['grayL'])
tx(s,'Balance T1：P_L = 16/19 = 0.8421 → R_L = 0.1579；P_R = 18/18 = 1 → R_R = 0',1.15,5.70,11.0,0.35,13,C['dark'],True,PP_ALIGN.CENTER)
takeaway(s,'T_d 是「被算了幾次」；U_d 是「真正有幾個不同 training objects」；P_d / R_d 描述重複程度。'); footer(s)

s=new_slide(); new.append(s)
title(s,'Step 3 — Standard 為什麼從 L 變成 R？','Balance Scale · Target 1：把公式完整算到 winner')
card(s,0.75,1.38,3.65,3.95,'Standard Voting','Score_STD(L) = T_L = 19\nScore_STD(R) = T_R = 18\n\n19 > 18\n\nSTD winner = L\nTrue = R\n\n→ WRONG',C['redL'],C['red'],15,13)
card(s,4.83,1.38,3.65,3.95,'Population Standard','Score_PS(L) = U_L = 16\nScore_PS(R) = U_R = 18\n\n16 < 18\n\nPS winner = R\nTrue = R\n\n→ CORRECT',C['greenL'],C['green'],15,13)
card(s,8.91,1.38,3.65,3.95,'到底改了什麼？','L：T=19 → U=16\n代表 3 次 support 是重複計價。\n\nR：T=18 → U=18\n沒有重複。\n\n所以不是「給重複資料負分」，而是把同一 training object 的重複 contribution 去重。',C['tealL'],C['teal'],14,11.3)
formula(s,'Score_STD(d|x) = T_d',1.0,5.62,3.25,0.58,16,C['grayL'])
formula(s,'Score_PS(d|x) = U_d',5.05,5.62,3.25,0.58,16,C['grayL'])
formula(s,'L → R  =  WRONG → CORRECTED',9.05,5.62,3.25,0.58,15,C['greenL'],C['green'])
takeaway(s,'這就是我們在 Target 回查裡標記 CORRECTED 的意思：同一筆資料只改 prediction policy，winner 從錯誤 class 換成 true class。'); footer(s)

s=new_slide(); new.append(s)
title(s,'Step 4 — GlobalStrength 到 Population Global 又怎麼算？','Car Evaluation · Target 2：原本 acc 勝，乘上 nonredundancy 後 good 勝')
formula(s,'GS_d = U_d / N_d',0.85,1.32,3.55,0.62,17,C['blueL'],C['blue'])
formula(s,'P_d = U_d / T_d',4.88,1.32,3.55,0.62,17,C['orangeL'],C['orange'])
formula(s,'PG_d = GS_d × P_d = U_d² / (N_d T_d)',8.55,1.32,3.95,0.62,16,C['tealL'],C['teal'])
card(s,0.75,2.25,5.8,2.45,'原 GlobalStrength','True = good\n\nacc: U=35, N=307 → GS=35/307=0.1140\ngood: U=5, N=55 → GS=5/55=0.0909\n\n0.1140 > 0.0909 → winner = acc → WRONG',C['redL'],C['red'],14,12)
card(s,6.8,2.25,5.8,2.45,'Population Global','acc: T=50, U=35 → P=35/50=0.70\nPG(acc)=0.1140×0.70=0.0798\n\ngood: T=5, U=5 → P=1\nPG(good)=0.0909×1=0.0909\n\n0.0909 > 0.0798 → winner = good → CORRECT',C['greenL'],C['green'],14,11.6)
card(s,2.15,5.03,9.05,0.95,'關鍵','GlobalStrength 已經使用 unique support；Population Global 額外問「這些 unique support 是不是由很多重疊 rules 反覆產生？」因此再乘 P_d。',C['tealL'],C['teal'],13,11.3)
takeaway(s,'PG 不是把 numerator / denominator 顛倒，而是在 GlobalStrength 上再加入 support-overlap penalty。'); footer(s)

s=new_slide(); new.append(s)
title(s,'從 15 個 Target 到 908 個 Conflict：怎麼擴大回查？','15 筆用來理解機制；908 筆則是把同一套計算對所有 conflict occurrences 全掃')
card(s,0.75,1.38,5.65,2.25,'15 個 selected targets','3 datasets × 5 cases\n\n目的：看清楚 corrected、harmed、winner 不變、true class 缺席等代表性機制。\n\n它們是 descriptive backtrace，不是 statistical sample。',C['blueL'],C['blue'],15,11.7)
card(s,6.85,1.38,5.65,2.25,'全體 conflict scan','對每個 dataset × C ∈ {2,3,4,5} 的 20% holdout：\n只要 |Candidate(x)| > 1，就記成一個 prediction-time conflict occurrence。\n\n同一 row 在不同 C 下可以重複出現。',C['tealL'],C['teal'],15,11.4)
formula(s,'Conflict(x) ⇔ | { d_r : r∈M(x) } | > 1',1.0,4.05,5.25,0.7,17,C['grayL'])
card(s,6.75,3.95,5.35,1.05,'實際計數','Balance: 47+62+74+83 = 266\nCar: 85+134+169+210 = 598\nHayes: 3+14+10+17 = 44',C['goldL'],C['gold'],13,11.8)
formula(s,'266 + 598 + 44 = 908 conflict occurrences',2.15,5.38,9.05,0.72,19,C['redL'],C['red'])
takeaway(s,'908 不是 908 個 unique people，而是 dataset × client-setting 下的 conflict occurrences。'); footer(s)

s=new_slide(); new.append(s)
title(s,'908 裡先分兩類：能不能被 voting 解決？','先看 true class 有沒有進 Candidate(x)，再談 ranking')
formula(s,'Resolvable(x) ⇔ y ∈ Candidate(x)',0.9,1.35,5.55,0.68,18,C['greenL'],C['green'])
formula(s,'Unresolvable(x) ⇔ y ∉ Candidate(x)',6.85,1.35,5.55,0.68,18,C['redL'],C['red'])
card(s,0.85,2.35,5.65,2.45,'Coverage / Induction Failure','例如 Balance Target 5：True = B\nCandidate(x) = {L, R}\n\nB 根本不存在，任何 FIRST / SIMPLE / STD / GS / PS / PG 都不可能選到 B。\n\n這種錯誤不能怪 voting。',C['redL'],C['red'],14,11.6)
card(s,6.85,2.35,5.65,2.45,'Conflict Ranking Problem','如果 y ∈ Candidate(x)，代表 true class 至少有 rule 提出。\n\n此時才有意義比較：\nT_d、U_d、P_d、GS_d、PG_d\n以及 old winner → new winner。',C['greenL'],C['green'],14,11.6)
rect(s,3.68,5.18,5.95,0.7,C['goldL']); tx(s,'908 conflicts 中有 45 筆：true class 不在 candidate set',3.82,5.39,5.65,0.24,13,C['gold'],True,PP_ALIGN.CENTER)
takeaway(s,'所以我們新增 Target Candidate Coverage，就是專門量「true class 是否有進 candidate set」。'); footer(s)

s=new_slide(); new.append(s)
title(s,'43 個 Winner Changes 到底怎麼算？','只要 baseline winner 與 population winner 不同，就記一次 transition')
formula(s,'Δ_STD(x) = 1[ ŷ_STD(x) ≠ ŷ_PS(x) ]',0.9,1.35,5.55,0.7,17,C['blueL'],C['blue'])
formula(s,'Δ_GS(x) = 1[ ŷ_GS(x) ≠ ŷ_PG(x) ]',6.85,1.35,5.55,0.7,17,C['tealL'],C['teal'])
card(s,0.85,2.38,5.65,2.55,'Standard → Population Standard','Winner changed = 16\n\n6  Corrected: wrong → true\n9  Harmed: true → wrong\n1  Changed, same correctness: wrong → another wrong',C['blueL'],C['blue'],14,12.2)
card(s,6.85,2.38,5.65,2.55,'Global → Population Global','Winner changed = 27\n\n10 Corrected: wrong → true\n13 Harmed: true → wrong\n4  Changed, same correctness',C['tealL'],C['teal'],14,12.2)
formula(s,'16 + 27 = 43 population-induced winner changes',2.05,5.32,9.25,0.68,18,C['goldL'],C['gold'])
takeaway(s,'43 不是「43 筆變好」；它只代表 Population correction 真正改變了最後 winner，之後還要分 corrected / harmed。'); footer(s)

s=new_slide(); new.append(s)
title(s,'Redundancy Asymmetry：最後這個量是怎麼算出來的？','不是另一個 prediction method；它是用來解釋「為什麼同一個 penalty 有時救對、有時修壞」的 diagnostic')
formula(s,'R_d = 1 - P_d = 1 - U_d/T_d',0.95,1.28,5.25,0.68,18,C['grayL'])
formula(s,'A(x) = R_competitor - R_true',7.05,1.28,5.25,0.68,18,C['grayL'])
card(s,0.78,2.25,5.75,2.68,'Corrected example · Car Target 2','True = good；原 GS winner = acc\n\nR_acc = 1 - 35/50 = 0.300\nR_good = 1 - 5/5 = 0\n\nA(x) = 0.300 - 0 = +0.300\n\n錯誤 competitor 比 true class 更重複 → penalty 壓低 acc → good 勝。',C['greenL'],C['green'],14,11.4)
card(s,6.82,2.25,5.75,2.68,'Harmed example · Car Target 4','True = acc；原 GS winner = acc，PG 改成 unacc\n\nR_unacc = 1 - 106/106 = 0\nR_acc = 1 - 39/67 = 0.4179\n\nA(x) = 0 - 0.4179 = -0.4179\n\ntrue class 反而更重複 → penalty 壓掉正確 evidence。',C['redL'],C['red'],14,11.15)
rect(s,2.2,5.25,8.95,0.62,C['tealL']); tx(s,'A(x) > 0：較像 corrected-side mechanism　｜　A(x) < 0：較像 over-correction-side mechanism',2.35,5.44,8.65,0.22,12.5,C['teal'],True,PP_ALIGN.CENTER)
takeaway(s,'A(x) 是 mechanism hypothesis / diagnostic；目前不能把它宣稱成已證明的 causal predictor。'); footer(s)

s=new_slide(); new.append(s)
title(s,'從一筆 Target 到研究結論：完整邏輯不要再跳步','這一頁把後半段所有數字放回同一條計算鏈')
chain=[('x','20% target',C['blueL'],C['blue']),('M(x)','matching rules',C['tealL'],C['teal']),('Candidate(x)','可選 decisions',C['goldL'],C['gold']),('S_G(r)','training supports',C['orangeL'],C['orange']),('T/U/P/R','重複結構',C['redL'],C['red']),('Scores','STD/PS/GS/PG',C['greenL'],C['green']),('Outcome','correct / harm / fail',C['grayL'],C['gray'])]
for i,(a,b,f,ac) in enumerate(chain):
    x=0.48+i*1.82; rect(s,x,1.6,1.5,1.05,f); tx(s,a,x+0.05,1.8,1.4,0.28,14,ac,True,PP_ALIGN.CENTER); tx(s,b,x+0.05,2.18,1.4,0.2,9.3,C['dark'],False,PP_ALIGN.CENTER)
    if i<len(chain)-1: line(s,x+1.52,2.12,x+1.78,2.12,C['line'],1.3)
card(s,0.72,3.15,3.75,2.15,'Case-level','15 selected targets\n→ 用來看清楚一筆資料為什麼翻轉\n→ 不是顯著性檢定',C['blueL'],C['blue'],14,11.5)
card(s,4.79,3.15,3.75,2.15,'Exhaustive conflict scan','908 conflict occurrences\n45 true class absent\n43 population-induced winner changes',C['goldL'],C['gold'],14,11.5)
card(s,8.86,3.15,3.75,2.15,'Mechanism hypothesis','43/43 changed winner → lower redundancy\n再用 A(x)=R_comp-R_true 解釋 corrected / harmed',C['tealL'],C['teal'],14,11.2)
formula(s,'Client Fragmentation → Rule Structure → Conflict → Support Redundancy → Redundancy Asymmetry → Prediction Policy → Performance',0.9,5.58,11.55,0.7,14,C['grayL'])
takeaway(s,'現在後半段不是「直接跳到 908 / 43」，而是每個統計數字都能回到單筆 target 的公式與 support objects。'); footer(s)

ids=list(prs.slides._sldIdLst)
orig_first15=ids[:15]
orig_targets=ids[15:18]
orig_tail=ids[18:23]
new_ids=ids[23:33]
order=orig_first15 + new_ids[:5] + orig_targets + new_ids[5:] + orig_tail
sld=prs.slides._sldIdLst
for el in list(sld): sld.remove(el)
for el in order: sld.append(el)

repls={
    '9. Balance Scale：Target 1–5 全部回查':'Target 回查案例 — Balance Scale · 1–5',
    '9. Car Evaluation：Target 1–5 全部回查':'Target 回查案例 — Car Evaluation · 1–5',
    '9. Hayes-Roth：Target 1–5 全部回查':'Target 回查案例 — Hayes-Roth · 1–5',
    '10. 908 個 Conflict Occurrences：把 error 拆開':'Summary — 908 個 Conflict Occurrences',
    '11. Redundancy Asymmetry：現在已正式寫進 code 與 CSV':'Summary — Redundancy Asymmetry',
    '12. 老師上次問題：現在逐題的最終答案':'老師上次問題：現在逐題的最終答案',
    '13. Code / Evidence Closure':'Code / Evidence Closure',
    '14. 現在可以下什麼結論、還不能下什麼結論':'現在可以下什麼結論、還不能下什麼結論',
}
for s in prs.slides:
    for sh in s.shapes:
        if hasattr(sh,'text') and sh.text in repls:
            sh.text_frame.paragraphs[0].runs[0].text=repls[sh.text]

for idx,s in enumerate(prs.slides,1):
    changed=False
    for sh in s.shapes:
        if not hasattr(sh,'text'): continue
        text=sh.text.strip()
        try: int(text)
        except: continue
        if sh.top > Inches(6.8) and sh.left > Inches(11.5):
            sh.text=str(idx); changed=True
    if not changed:
        for sh in s.shapes:
            if hasattr(sh,'text') and sh.text.strip()=='0' and sh.top > Inches(6.8): sh.text=str(idx)

prs.save(OUT)
OUT.replace(SRC)
print(SRC)
print('slides',len(prs.slides))
