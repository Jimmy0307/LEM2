from __future__ import annotations
import argparse, json, hashlib
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY=RGBColor(0x1F,0x29,0x33); DARK=RGBColor(0x23,0x31,0x3F); TEAL=RGBColor(0x2B,0x7A,0x78)
BLUE=RGBColor(0x4A,0x78,0xA8); GREEN=RGBColor(0x4C,0x8A,0x61); ORANGE=RGBColor(0xC9,0x6A,0x3D)
GOLD=RGBColor(0xA9,0x80,0x25); RED=RGBColor(0xB6,0x52,0x52); GRAY=RGBColor(0x6B,0x77,0x80)
LIGHT2=RGBColor(0xEB,0xF1,0xF3); WHITE=RGBColor(0xFF,0xFF,0xFF)
PALE_TEAL=RGBColor(0xE8,0xF3,0xF1); PALE_BLUE=RGBColor(0xE9,0xF0,0xF7); PALE_ORANGE=RGBColor(0xF8,0xEE,0xE9); PALE_RED=RGBColor(0xF8,0xE9,0xE9)
FONT='Noto Sans'

FROZEN_CONFIG = {
    'Balance Scale': {2:(47,37.3),3:(62,49.2),4:(74,58.7),5:(83,65.9)},
    'Car Evaluation': {2:(85,24.6),3:(134,38.7),4:(169,48.8),5:(210,60.7)},
    'Hayes-Roth': {2:(3,11.5),3:(14,53.8),4:(10,38.5),5:(17,65.4)},
}
DATASET_BEST = {
    'Balance Scale': ('PS',195,266,73.3),
    'Car Evaluation': ('FIRST',486,598,81.3),
    'Hayes-Roth': ('FIRST',36,44,81.8),
}

def add_text(slide,text,x,y,w,h,size=11,bold=False,color=NAVY,align=PP_ALIGN.LEFT,valign=MSO_ANCHOR.TOP):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame
    tf.clear(); tf.word_wrap=True; tf.vertical_anchor=valign; tf.margin_left=tf.margin_right=Inches(.03); tf.margin_top=tf.margin_bottom=Inches(.02)
    p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=str(text); r.font.name=FONT; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    return box

def replace_text(shape,text,size=None,color=None,bold=None):
    shape.text=str(text); shape.text_frame.word_wrap=True
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.name=FONT
            if size is not None: r.font.size=Pt(size)
            if color is not None: r.font.color.rgb=color
            if bold is not None: r.font.bold=bold

def add_rect(slide,x,y,w,h,fill=WHITE,line=LIGHT2,radius=True,line_width=1):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=fill; sh.line.color.rgb=line; sh.line.width=Pt(line_width); return sh

def add_title(slide,title,subtitle):
    add_text(slide,title,.62,.34,12.05,.48,23.5,True,NAVY); add_text(slide,subtitle,.64,.86,12,.34,12,False,TEAL)

def add_footer(slide,page):
    add_text(slide,'Federated LEM2 · Conflict concentration audit · 2026-10-25',.62,7.13,10.5,.18,7.8,False,GRAY)
    add_text(slide,str(page),12.2,7.13,.4,.18,7.8,False,GRAY,PP_ALIGN.RIGHT)

def add_conclusion(slide,text):
    add_text(slide,'結論：',.92,6.43,.8,.22,9.5,True,TEAL); add_text(slide,text,1.65,6.39,10.55,.32,9.4,False,NAVY)

def add_kpi(slide,x,y,w,h,value,label,accent=TEAL,note=None):
    add_rect(slide,x,y,w,h,WHITE,LIGHT2); stripe=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(.08),Inches(h)); stripe.fill.solid(); stripe.fill.fore_color.rgb=accent; stripe.line.fill.background()
    add_text(slide,value,x+.22,y+.16,w-.32,.38,22,True,accent); add_text(slide,label,x+.22,y+.60,w-.32,.32,10.2,True,NAVY)
    if note: add_text(slide,note,x+.22,y+.96,w-.32,h-1.02,8.8,False,GRAY)

def set_cell(cell,text,size=10,bold=False,color=NAVY,fill=WHITE,align=PP_ALIGN.CENTER):
    cell.fill.solid(); cell.fill.fore_color.rgb=fill; cell.text=''; tf=cell.text_frame; tf.clear(); tf.vertical_anchor=MSO_ANCHOR.MIDDLE; tf.margin_left=tf.margin_right=Inches(.04); tf.margin_top=tf.margin_bottom=Inches(.02)
    p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=str(text); r.font.name=FONT; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color

def title_slide(prs,title):
    for slide in prs.slides:
        for sh in slide.shapes:
            if hasattr(sh,'text') and sh.text.strip()==title: return slide
    raise KeyError(title)

def add_concentration_slide(prs, summary):
    s=prs.slides.add_slide(prs.slide_layouts[6]); add_title(s,'908 個 Conflict 到底集中在哪裡？','不是平均分布：Car 佔 65.9%；高 client-count configurations 承擔更大的 conflict workload')
    rows=[['Dataset','C=2','C=3','C=4','C=5','Total']]
    names=[('Balance Scale','Balance'),('Car Evaluation','Car'),('Hayes-Roth','Hayes')]
    for full,short in names:
        vals=[short]
        for c in [2,3,4,5]: n,r=FROZEN_CONFIG[full][c]; vals.append(f'{n}\n{r:.1f}%')
        vals.append(str(summary['dataset_conflicts'][full])); rows.append(vals)
    tab=s.shapes.add_table(4,6,Inches(.75),Inches(1.55),Inches(8),Inches(3.9)).table
    for i,w in enumerate([1.45,1.15,1.15,1.15,1.15,1.35]): tab.columns[i].width=Inches(w)
    for r,row in enumerate(rows):
        tab.rows[r].height=Inches(.66 if r==0 else 1.08)
        for c,val in enumerate(row):
            if r==0: set_cell(tab.cell(r,c),val,10.2,True,WHITE,DARK)
            elif c==0: set_cell(tab.cell(r,c),val,10.5,True,NAVY,LIGHT2,PP_ALIGN.LEFT)
            elif c==5: set_cell(tab.cell(r,c),val,13.5,True,TEAL,PALE_TEAL)
            else:
                rate=float(val.split('\n')[1].strip('%')); fill=PALE_BLUE if rate<40 else (PALE_TEAL if rate<55 else PALE_ORANGE); set_cell(tab.cell(r,c),val,10.5,True,NAVY,fill)
    add_text(s,'每格：conflict count / configuration-level conflict rate',.82,5.56,7.7,.28,8.8,False,GRAY)
    add_kpi(s,9.05,1.55,3.4,1.28,'598 / 908','Car = 65.9% of all conflicts',BLUE,'workload 高度集中在 Car Evaluation')
    add_kpi(s,9.05,3.02,3.4,1.28,'310 / 908','C=5 = 34.1%',ORANGE,'只看 C=5 就占全部 occurrences 三分之一')
    add_kpi(s,9.05,4.49,3.4,1.28,'908 / 1,992','Overall workload = 45.6%',TEAL,'configuration-level test evaluations')
    add_conclusion(s,'908 不是一個均質群體；先按 dataset × C 分解，才能判斷真正的 conflict workload 集中點。'); return s

def add_persistence_slide(prs,summary):
    s=prs.slides.add_slide(prs.slide_layouts[6]); add_title(s,'908 次其實只有 392 個 unique objects：Conflict 會反覆出現','同一 (dataset, fold, row_id) 可在不同 C 再次 conflict；occurrence 與 unique object 必須分開')
    dist=summary['persistence_distribution']; cards=[('1 個 C',str(dist['1']),'SINGLE',BLUE),('2 個 C',str(dist['2']),'REPEATED',TEAL),('3 個 C',str(dist['3']),'REPEATED',GOLD),('4 個 C',str(dist['4']),'PERSISTENT',ORANGE)]
    for i,(lab,val,status,col) in enumerate(cards):
        x=.82+i*3.05; add_rect(s,x,1.62,2.6,1.52,WHITE,LIGHT2); add_text(s,val,x+.12,1.82,2.36,.42,24,True,col,PP_ALIGN.CENTER); add_text(s,'unique objects',x+.12,2.28,2.36,.26,9.5,False,GRAY,PP_ALIGN.CENTER); add_text(s,f'{lab} · {status}',x+.12,2.63,2.36,.28,9.5,True,NAVY,PP_ALIGN.CENTER)
    add_rect(s,.92,3.53,7.25,1.55,PALE_BLUE,LIGHT2); add_text(s,'106×1 + 123×2 + 96×3 + 67×4 = 908 occurrences',1.15,3.82,6.8,.38,16,True,NAVY,PP_ALIGN.CENTER); add_text(s,'Persistence(x) = x 在 C={2,3,4,5} 中被判為 conflict 的 distinct client-count 數',1.18,4.35,6.75,.35,10,False,GRAY,PP_ALIGN.CENTER)
    add_kpi(s,8.55,3.42,3.75,1.02,'286 / 392 = 73.0%','至少在 2 個 C 重複 conflict',TEAL); add_kpi(s,8.55,4.58,3.75,1.02,'67 / 392 = 17.1%','四種 C 全部 conflict',ORANGE)
    add_text(s,'67 個 persistent objects 貢獻 268 / 908 = 29.5% 的 conflict occurrences。',1.02,5.42,7,.4,10.5,True,TEAL); add_text(s,'解讀邊界：這證明 frozen split 下存在「反覆出現」的 subset；它提示 structural overlap，但本身還不是 causal proof。',1.02,5.84,11.25,.38,9.2,False,GRAY)
    add_conclusion(s,'Conflict 同時包含 partition-sensitive 與 persistent cases；不能把 908 全部直接叫成 structural conflict。'); return s

def add_boundary_slide(prs,summary):
    s=prs.slides.add_slide(prs.slide_layouts[6]); add_title(s,'Conflict 不是平均撞：高度集中在少數 decision boundaries','Candidate(x) 組合顯示 cross-decision overlap 有明顯結構，不是在所有 class pairs 均勻散開')
    cfg=[('Balance Scale',TEAL,'R 128 · L 106 · B 32'),('Car Evaluation',BLUE,'unacc 374 · acc 151 · good 41 · vgood 32'),('Hayes-Roth',ORANGE,'class1 18 · class3 14 · class2 12')]
    totals=summary['dataset_conflicts']
    for i,(ds,col,truecounts) in enumerate(cfg):
        top=summary['top_candidate_signatures'][ds][0]; sig=top['candidate_signature'].replace('|',' | '); n=top['conflict_count']; pct=100*n/totals[ds]; x=.78+i*4.15
        add_rect(s,x,1.56,3.72,3.78,WHITE,LIGHT2); add_text(s,ds,x+.22,1.78,3.28,.30,12.2,True,NAVY,PP_ALIGN.CENTER); add_text(s,'Top candidate signature',x+.22,2.28,3.28,.24,8.7,False,GRAY,PP_ALIGN.CENTER); add_text(s,sig,x+.22,2.62,3.28,.50,20,True,col,PP_ALIGN.CENTER); add_text(s,f'{n} / {totals[ds]} = {pct:.1f}%',x+.22,3.21,3.28,.34,12.2,True,NAVY,PP_ALIGN.CENTER)
        add_rect(s,x+.42,3.72,2.88,.18,LIGHT2,LIGHT2,False,0); bar=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x+.42),Inches(3.72),Inches(2.88*pct/100),Inches(.18)); bar.fill.solid(); bar.fill.fore_color.rgb=col; bar.line.fill.background(); add_text(s,'True-class conflict occurrences',x+.22,4.20,3.28,.23,8.4,False,GRAY,PP_ALIGN.CENTER); add_text(s,truecounts,x+.22,4.52,3.28,.55,9.4,True,NAVY,PP_ALIGN.CENTER)
    add_rect(s,.92,5.58,11.42,.58,PALE_RED,PALE_RED); add_text(s,'注意：true-class / factor-value 目前都是「conflict 內部的出現次數」，沒有 non-conflict denominator，因此不能直接解讀成 class risk 或因果效果。',1.10,5.72,11.05,.28,9,False,RED,PP_ALIGN.CENTER); add_conclusion(s,'最主要的 conflict boundary 已可定位：Balance 的 L↔R、Car 的 acc↔unacc、Hayes 的 class 1↔2 最集中。'); return s

def add_method_issues_slide(prs):
    s=prs.slides.add_slide(prs.slide_layouts[6]); add_title(s,'同一個 Conflict，六種方法各自可能在哪裡出問題？','Conflict set 完全相同；差別不是「誰製造 conflict」，而是 resolver 使用哪種 evidence、因此產生哪種 bias')
    items=[('FIRST_MATCH','Order bias：第一條 matching rule 直接決定 winner。\n較晚出現、即使證據更強的 rule 也可能完全被忽略。',PALE_ORANGE,ORANGE),('SIMPLE','Rule-count bias：每條 rule 等權。\n多條寬／重疊 rules 可用票數壓過較少但 evidence 更強的 rules。',RGBColor(0xF8,0xF3,0xDE),GOLD),('STANDARD','Duplication bias：比較 T_d = support occurrences。\n同一 training object 可跨多條 matching rules 被重複計價。',PALE_BLUE,BLUE),('GLOBAL','Relative-coverage sensitivity：比較 U_d / N_d。\n已對 object 去重，但看的是 class 內「相對覆蓋率」，不是絕對 evidence 量。',RGBColor(0xE6,0xF1,0xE7),GREEN),('POPULATION STANDARD','De-dup over-correction：只比較 unique U_d。\n16 次翻 winner：6 corrected · 9 harmed · 1 wrong→wrong。',PALE_TEAL,TEAL),('POPULATION GLOBAL','Redundancy-penalty over-correction：GS × U_d/T_d。\n27 次翻 winner：10 corrected · 13 harmed · 4 same correctness。',PALE_RED,RED)]
    for i,(name,body,fill,col) in enumerate(items):
        row=i//2; colidx=i%2; x=1.0+colidx*6.05; y=1.45+row*1.48; add_rect(s,x,y,5.65,1.16,fill,LIGHT2); stripe=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x+.12),Inches(y+.16),Inches(.07),Inches(.87)); stripe.fill.solid(); stripe.fill.fore_color.rgb=col; stripe.line.fill.background(); add_text(s,name,x+.34,y+.18,5.0,.28,11.2,True,NAVY); add_text(s,body,x+.34,y+.56,5.0,.49,9.4,False,NAVY)
    add_text(s,'Mechanism-level failure modes',1.0,5.94,4.8,.28,9.2,True,GRAY,PP_ALIGN.CENTER); add_text(s,'Observed population-correction outcomes',7.1,5.94,4.8,.28,9.2,True,RED,PP_ALIGN.CENTER); add_conclusion(s,'同一 908 cases 的 resolver audit 已完成：整體 conflict-only accuracy 最高為 STANDARD 77.0%，但各 dataset 的最佳 resolver 不同。'); return s

def add_resolver_slide(prs,summary):
    s=prs.slides.add_slide(prs.slide_layouts[6]); add_title(s,'同一 908 個 Conflict：六種 Resolver 到底誰比較會選對？','Conflict-only accuracy；六種方法的分母完全相同 = 908，現在可以直接比較 resolver 本身')
    order=[('STANDARD_VOTING','STANDARD',TEAL),('POPULATION_STANDARD','POPULATION STANDARD',GREEN),('GLOBAL_STRENGTH','GLOBAL STRENGTH',BLUE),('POPULATION_GLOBAL','POPULATION GLOBAL',GOLD),('FIRST_MATCH','FIRST MATCH',ORANGE),('SIMPLE_VOTING','SIMPLE VOTING',GRAY)]
    add_rect(s,.72,1.47,7.52,4.72,WHITE,LIGHT2)
    for i,(key,label,col) in enumerate(order):
        rec=summary['policy_overall'][key]; acc=rec['conflict_accuracy']; correct=rec['correct']; y=1.78+i*.68; add_text(s,label,.98,y,2.2,.28,9.1,True,NAVY); bg=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(3.08),Inches(y+.02),Inches(3.65),Inches(.22)); bg.fill.solid(); bg.fill.fore_color.rgb=LIGHT2; bg.line.fill.background(); fg=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(3.08),Inches(y+.02),Inches(3.65*acc),Inches(.22)); fg.fill.solid(); fg.fill.fore_color.rgb=col; fg.line.fill.background(); add_text(s,f'{acc*100:.1f}%  ({correct}/908)',6.92,y-.03,1.05,.30,9,True,col,PP_ALIGN.RIGHT)
    add_rect(s,8.52,1.47,3.82,2.35,PALE_BLUE,LIGHT2); add_text(s,'Dataset-specific best resolver',8.78,1.72,3.30,.30,11.2,True,NAVY,PP_ALIGN.CENTER); add_text(s,'Balance   PS  195/266 = 73.3%\nCar       FIRST 486/598 = 81.3%\nHayes     FIRST 36/44 = 81.8%',8.90,2.20,3.05,1.10,10.1,True,NAVY); add_text(s,'→ 沒有 universal winner',8.80,3.33,3.25,.25,9.8,True,TEAL,PP_ALIGN.CENTER)
    t=summary['transition_totals']; add_rect(s,8.52,4.03,3.82,2.08,PALE_ORANGE,LIGHT2); add_text(s,'Population correction 的 net effect',8.78,4.25,3.30,.30,10.8,True,NAVY,PP_ALIGN.CENTER); add_text(s,f"STD → PS\n{t['STD_TO_PS']['CORRECTED']} corrected vs {t['STD_TO_PS']['HARMED']} harmed → net −3 correct\n\nGS → PG\n{t['GS_TO_PG']['CORRECTED']} corrected vs {t['GS_TO_PG']['HARMED']} harmed → net −3 correct",8.90,4.72,3.05,1.20,9.2,False,NAVY); add_conclusion(s,'Population correction 的行為符合去重設計，但在目前 frozen experiment 並沒有整體勝過原 Standard / Global baseline。'); return s

def build(source:Path,summary_path:Path,output:Path):
    summary=json.loads(summary_path.read_text(encoding='utf-8')); assert summary['total_conflict_occurrences']==908 and summary['unique_conflict_objects']==392
    prs=Presentation(source)
    # Refuse to double-add.
    for t in ['908 個 Conflict 到底集中在哪裡？','同一個 Conflict，六種方法各自可能在哪裡出問題？']:
        try: title_slide(prs,t); raise RuntimeError(f'source already contains generated slide: {t}')
        except KeyError: pass
    # Clarify the existing C→Conflict slide into direct vs structural causes.
    cause = title_slide(prs,'為什麼 C 從 2 增加到 5，Conflict 通常會上升？')
    for sh in cause.shapes:
        if not hasattr(sh,'text'): continue
        if sh.text.strip()=='為什麼 C 從 2 增加到 5，Conflict 通常會上升？': replace_text(sh,'Conflict 為什麼會產生？直接原因 vs. 結構性原因',23.5,NAVY,True)
        elif sh.text.startswith('目前資料最一致的結構鏈'): replace_text(sh,'直接原因是 cross-decision rule overlap；client fragmentation 只是在目前資料中提高 overlap 的機率',12,TEAL)
        elif sh.text.strip()=='直覺': replace_text(sh,'結構性原因',11,TEAL,True)
        elif sh.text.startswith('同樣 80% training'): replace_text(sh,'同樣 80% training 切成更多 clients 後，每個 client 可見資料變少。\nLocal LEM2 在目前資料中傾向形成較短、較一般的 rules；UNION 後更容易同時覆蓋同一個 test x。',9.5,NAVY)
        elif sh.text.strip()=='但不是定理': replace_text(sh,'直接原因',11,TEAL,True)
        elif sh.text.startswith('Hayes test 只有'): replace_text(sh,'真正的 Conflict 條件永遠是 |Candidate(x)| > 1：同一個 x 同時命中「不同 decision」的 rules。\nFragmentation 本身不等於 conflict；Hayes C=3→4 的回落也顯示它不是單調定理。',9.5,NAVY)
        elif sh.text.startswith('Conflict 上升是 rule structure'): replace_text(sh,'直接原因 = 同一 x 同時命中不同 decision rules；結構原因 = fragmentation 改變 rule generality / overlap。',9.2,NAVY)
    sA=add_concentration_slide(prs,summary); sB=add_persistence_slide(prs,summary); sC=add_boundary_slide(prs,summary); sM=add_method_issues_slide(prs); sD=add_resolver_slide(prs,summary); new=[sA,sB,sC,sM,sD]
    # existing edits
    for sh in title_slide(prs,'數據真的有沿著這條鏈走嗎？Car / Balance：有').shapes:
        if hasattr(sh,'text') and sh.text.startswith('這就是我們目前說'): replace_text(sh,'在 Car / Balance 的 frozen sweep 中，client fragmentation、rule generality、matching overlap 與 conflict workload 呈一致方向變化。',9.2,NAVY)
    for sh in title_slide(prs,'六種方法到底改了什麼？它們只改「Conflict 發生後的排序」').shapes:
        if hasattr(sh,'text') and sh.text.strip()=='908 = structural conflict occurrences': replace_text(sh,'908 = configuration-level conflict occurrences',11,TEAL,True)
    for sh in title_slide(prs,'老師上次問題：現在逐題的最終答案').shapes:
        if hasattr(sh,'text') and sh.text.startswith('Coverage 與 Conflict 都已改成'): replace_text(sh,'Coverage / Conflict 均已有可重現定義；908 occurrences 已進一步拆成 392 unique objects，且 conflict 高度集中在少數 candidate boundaries。',9,NAVY)
    for sh in title_slide(prs,'Code / Evidence Closure').shapes:
        if not hasattr(sh,'text'): continue
        if sh.text.startswith('✓ Target Candidate Coverage'): replace_text(sh,'✓ Target Candidate Coverage\n✓ Redundancy Asymmetry helper\n✓ Conflict concentration / persistence audit\n✓ Candidate-signature / resolver audit\n✓ Standard / Global transition decomposition',9.2,NAVY)
        elif sh.text.startswith('✓ dataset metadata'): replace_text(sh,'✓ dataset metadata / six-policy metrics\n✓ coverage by class / target 1–5\n✓ conflict anatomy bundle (10 files)\n✓ persistence / candidate / policy tables\n✓ error / asymmetry transitions',9.2,NAVY)
        elif sh.text.startswith('TDD RED：'): replace_text(sh,'TDD RED：缺 function / entrypoint 時 CI fail\nGREEN：implementation 後\n124 tests passed · compile PASS\nConflict reconciliation PASS\nEvidence workflow SUCCESS',9.2,NAVY)
        elif sh.text.startswith('這一輪完成的是 professor-question evidence layer'): replace_text(sh,'這一輪再補上 Conflict Concentration / Anatomy evidence layer；仍不修改 frozen main，也不新增第七種 prediction policy。',9.1,NAVY)
        elif sh.text.startswith('Federated LEM2 · Professor-question closure'): replace_text(sh,'Federated LEM2 · Conflict concentration closure · 2026-10-25',7.8,GRAY)
    for sh in title_slide(prs,'現在可以下什麼結論、還不能下什麼結論').shapes:
        if not hasattr(sh,'text'): continue
        if sh.text.startswith('• Population correction 會降低'): replace_text(sh,'• 908 occurrences 來自 392 unique objects；67 個在四種 C 都 conflict。\n• Candidate conflict 高度集中：Balance L|R 56%、Car acc|unacc 56%、Hayes 1|2 68%。\n• 同一 908 cases：STD 77.0%、PS 76.7%、GS 72.8%、PG 72.5%；沒有 universal resolver。\n• 43/43 changed winners 都往 lower-redundancy decision，但可 corrected 也可 harmed。\n• true class 不在 Candidate(x) 時，任何 voting 都救不了。',9.1,NAVY)
        elif sh.text.startswith('• Population method universally better'): replace_text(sh,'• Population method universally better。\n• Client 數越多，conflict 就必然單調增加。\n• Persistent conflict 已被證明是 federation 的 causal effect。\n• Redundancy 本身就是 noise，或 A(x) 已是 causal predictor。\n• conflict 內 factor/value 頻率可直接解讀成 relative risk。\n• seed=42 的單次 split 可直接當獨立樣本做顯著性。',8.9,NAVY)
        elif sh.text.startswith('Repeated matched seeds / splits'): sh.top=Inches(5.76); sh.height=Inches(.45); sh.rotation=0; replace_text(sh,'Repeated matched seeds / splits → 追蹤 conflict persistence、candidate boundary、T/U/P、score margin、A(x) → block-level Friedman / paired Wilcoxon。',9,NAVY)
        elif sh.text.startswith('下一輪不加新方法'): replace_text(sh,'下一輪不加新方法；先驗證哪些 decision boundary / support structure 形成 persistent conflict，以及何時 Population correction corrected / over-correct。',9,NAVY)
        elif sh.text.startswith('Federated LEM2 · Professor-question closure'): replace_text(sh,'Federated LEM2 · Conflict concentration closure · 2026-10-25',7.8,GRAY)
    # reorder: A/B/C after 908; M after cause; D after M.
    lst=prs.slides._sldIdLst; elems={int(el.id):el for el in list(lst)}; ids=[x.slide_id for x in new]
    for sid in ids: lst.remove(elems[int(sid)])
    anchor=elems[int(title_slide(prs,'908 怎麼來？為什麼看起來這麼大？').slide_id)]; idx=list(lst).index(anchor)+1
    for sid in ids[:3]: lst.insert(idx,elems[int(sid)]); idx+=1
    anchor=elems[int(title_slide(prs,'Conflict 為什麼會產生？直接原因 vs. 結構性原因').slide_id)]; idx=list(lst).index(anchor)+1; lst.insert(idx,elems[int(ids[3])]); idx+=1; lst.insert(idx,elems[int(ids[4])])
    new_ids=set(ids)
    for i,slide in enumerate(prs.slides,1):
        if slide.slide_id in new_ids: add_footer(slide,i)
        else:
            found=False
            for sh in slide.shapes:
                if hasattr(sh,'text') and sh.text.strip().isdigit() and sh.top/914400>6.9 and sh.left/914400>11.7: replace_text(sh,str(i),7.8,GRAY); found=True
            if not found: add_text(slide,str(i),12.2,7.13,.4,.18,7.8,False,GRAY,PP_ALIGN.RIGHT)
    assert len(prs.slides)==38
    output.parent.mkdir(parents=True,exist_ok=True); prs.save(output)
    digest=hashlib.sha256(output.read_bytes()).hexdigest(); output.with_suffix(output.suffix+'.sha256').write_text(f'{digest}  {output.name}\n',encoding='utf-8')
    print(f'DECK_BUILD=PASS slides={len(prs.slides)} sha256={digest}')

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--source',type=Path,required=True); ap.add_argument('--summary',type=Path,required=True); ap.add_argument('--output',type=Path,required=True); a=ap.parse_args(); build(a.source,a.summary,a.output)
if __name__=='__main__': main()
