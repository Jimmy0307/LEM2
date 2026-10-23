from __future__ import annotations

from pathlib import Path
import pandas as pd
import numpy as np
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor

ROOT = Path('docs/research/conflict_client_sweep_20261020/evidence/20261023/teacher')
OUT = Path('docs/research/conflict_client_sweep_20261020/presentation')
OUT.mkdir(parents=True, exist_ok=True)
PPTX = OUT / 'Federated_LEM2_Teacher_Questions_Complete_20261023.pptx'

meta = pd.read_csv(ROOT/'TEACHER_DATASET_METADATA.csv')
metrics = pd.read_csv(ROOT/'TEACHER_SIX_POLICY_METRICS.csv')
coverage = pd.read_csv(ROOT/'TEACHER_COVERAGE_BY_CLASS.csv')
targets = pd.read_csv(ROOT/'TEACHER_TARGET_1_TO_5.csv')
stdtrans = pd.read_csv(ROOT/'TEACHER_STANDARD_TRANSITIONS.csv')
globtrans = pd.read_csv(ROOT/'TEACHER_GLOBAL_TRANSITIONS.csv')
errors = pd.read_csv(ROOT/'TEACHER_ERROR_DECOMPOSITION.csv')
asym = pd.read_csv(ROOT/'TEACHER_REDUNDANCY_ASYMMETRY.csv')

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
FONT = 'Noto Sans CJK TC'
C = {
    'dark':'1F2933','teal':'2B7A78','tealL':'E7F2F1','blue':'4A78A8','blueL':'EAF2F8',
    'orange':'C96A3D','orangeL':'F8EAE3','green':'4C8A61','greenL':'E8F3E8',
    'red':'B65252','redL':'FBEAEA','gray':'6B7780','grayL':'F5F7F8','line':'D9E2E7',
    'gold':'A98025','goldL':'F6F0DC','white':'FFFFFF'
}
def rgb(h): return RGBColor.from_string(h)
def slide():
    s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(C['white']); return s
def tx(s,text,x,y,w,h,size=14,color='1F2933',bold=False,align=PP_ALIGN.LEFT,valign=MSO_ANCHOR.TOP,font=FONT):
    box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0; tf.vertical_anchor=valign
    p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=str(text); r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=rgb(color); return box
def rect(s,x,y,w,h,fill,line=None,round=True):
    shp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); shp.fill.solid(); shp.fill.fore_color.rgb=rgb(fill); shp.line.color.rgb=rgb(line or fill); return shp
def line(s,x1,y1,x2,y2,color='D9E2E7',width=1):
    ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); ln.line.color.rgb=rgb(color); ln.line.width=Pt(width); return ln
def title(s,t,sub=None):
    tx(s,t,0.62,0.34,12.05,0.48,24,C['dark'],True)
    if sub: tx(s,sub,0.64,0.86,12.0,0.34,12.5,C['teal'])
def footer(s,n):
    line(s,0.55,7.07,12.75,7.07,C['line'],0.7); tx(s,'Federated LEM2 · Professor-question closure · 2026-10-23',0.62,7.13,10.5,0.18,8,C['gray']); tx(s,str(n),12.2,7.13,0.4,0.18,8,C['gray'],False,PP_ALIGN.RIGHT)
def card(s,x,y,w,h,head,body,fill='F5F7F8',accent='2B7A78',hs=14,bs=11.5):
    rect(s,x,y,w,h,fill); rect(s,x+0.12,y+0.15,0.055,h-0.3,accent,accent,False); tx(s,head,x+0.32,y+0.15,w-0.48,0.30,hs,C['dark'],True); tx(s,body,x+0.32,y+0.56,w-0.50,h-0.66,bs,C['dark'])
def takeaway(s,text):
    rect(s,0.72,6.31,11.9,0.56,C['tealL']); tx(s,'結論：',0.92,6.46,0.8,0.2,10.5,C['teal'],True); tx(s,text,1.65,6.42,10.55,0.26,10.7,C['dark'])
def add_table(s, df, x,y,w,h, col_widths=None, header_fill='E7F2F1', font_size=10.5, first_col_bold=False, formats=None):
    rows, cols = df.shape[0]+1, df.shape[1]; table=s.shapes.add_table(rows,cols,Inches(x),Inches(y),Inches(w),Inches(h)).table
    if col_widths:
        total=sum(col_widths)
        for i,cw in enumerate(col_widths): table.columns[i].width=Inches(w*cw/total)
    for j,col in enumerate(df.columns):
        cell=table.cell(0,j); cell.text=str(col); cell.fill.solid(); cell.fill.fore_color.rgb=rgb(header_fill); p=cell.text_frame.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
        for r in p.runs: r.font.name=FONT; r.font.size=Pt(font_size); r.font.bold=True; r.font.color.rgb=rgb(C['dark'])
        cell.vertical_anchor=MSO_ANCHOR.MIDDLE
    for i in range(df.shape[0]):
        for j,col in enumerate(df.columns):
            v=df.iloc[i,j]
            if formats and col in formats: text=formats[col](v)
            elif isinstance(v,(float,np.floating)) and not pd.isna(v): text=f'{v:.4f}'
            else: text='' if pd.isna(v) else str(v)
            cell=table.cell(i+1,j); cell.text=text
            if i%2==1: cell.fill.solid(); cell.fill.fore_color.rgb=rgb('FAFBFC')
            p=cell.text_frame.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
            for r in p.runs: r.font.name=FONT; r.font.size=Pt(font_size); r.font.color.rgb=rgb(C['dark']); r.font.bold=(first_col_bold and j==0)
            cell.vertical_anchor=MSO_ANCHOR.MIDDLE
    return table
def fmt_pct(v): return '' if pd.isna(v) else f'{float(v)*100:.1f}%'
def short_policy(p): return {'FIRST_MATCH':'FIRST','SIMPLE_VOTING':'SIMPLE','STANDARD_VOTING':'STD','GLOBAL_STRENGTH':'GS','POPULATION_STANDARD':'PS','POPULATION_GLOBAL':'PG'}.get(p,p)

s=slide(); tx(s,'PROFESSOR QUESTION CLOSURE',0.65,0.48,5.5,0.28,12,C['gray'],True); tx(s,'Federated LEM2：把上次老師的問題逐題回答完整',0.65,1.18,11.9,0.72,29,C['dark'],True); tx(s,'資料規模 → 80/20 → 六方法 → Coverage → Standard / Global 改前改後 → 3×5 Target 回查 → 錯誤分解',0.68,2.25,11.7,0.65,16,C['teal'])
for i,(big,label,fill,accent) in enumerate([('3','UCI datasets',C['blueL'],C['blue']),('6','prediction methods',C['greenL'],C['green']),('15','target backtraces',C['orangeL'],C['orange']),('908','conflict occurrences',C['redL'],C['red'])]):
    x=0.85+i*3.02; rect(s,x,3.85,2.6,1.35,fill); tx(s,big,x+0.1,4.05,2.4,0.45,28,accent,True,PP_ALIGN.CENTER); tx(s,label,x+0.12,4.58,2.36,0.30,11,C['dark'],False,PP_ALIGN.CENTER)
tx(s,'這一版不先談「新發現」，而是先把老師上次每一題用 code → evidence → result 回答完。',1.0,5.72,11.3,0.45,14,C['dark'],True,PP_ALIGN.CENTER); footer(s,1)

s=slide(); title(s,'老師上次到底問了什麼？','本版簡報的順序完全照這張 checklist')
questions=[('1','三份資料有幾筆？幾個 headers / factors？'),('2','80/20 怎麼切？20% 有沒有參與建模？'),('3','同一資料跑六種方法，Acc / Recall / F1 到底是多少？'),('4','Overall Coverage、Target Coverage 到底是多少？'),('5','Standard 的 3+3 為什麼有問題？改完 A/B 誰贏？'),('6','Global 改法和原 GlobalStrength 差在哪？winner 怎麼變？'),('7','每個 dataset 的 Target 1–5 回查後，差別到底在哪？')]
for i,(n,q) in enumerate(questions):
    y=1.42+i*0.68; rect(s,0.85,y,0.46,0.46,C['tealL']); tx(s,n,0.85,y+0.08,0.46,0.2,12,C['teal'],True,PP_ALIGN.CENTER); tx(s,q,1.52,y+0.07,10.8,0.28,13,C['dark'],True if i in (4,5,6) else False)
takeaway(s,'後面的 Redundancy Asymmetry 只能放在「回答完老師」之後，不能反過來取代老師的問題。'); footer(s,2)

s=slide(); title(s,'1. 資料規模：Rows、Headers、Factors 全部明列','Headers = condition attributes + 1 decision column')
md=meta.sort_values(['dataset','client_count']).drop_duplicates('dataset')[['dataset','dataset_rows','dataset_columns','condition_attribute_count','target_column','class_count','train_rows','test_rows']]; md=md.rename(columns={'dataset':'Dataset','dataset_rows':'Rows','dataset_columns':'Headers','condition_attribute_count':'Factors','target_column':'Decision','class_count':'Classes','train_rows':'Train 80%','test_rows':'Test 20%'})
add_table(s,md,0.7,1.55,11.95,2.55,[2.1,1,1,1,1.1,1,1.1,1.1],font_size=11.5,first_col_bold=True); card(s,0.85,4.45,3.55,1.35,'Hayes-Roth','132 rows · 5 headers\n4 factors + 1 decision · 3 classes',C['blueL'],C['blue'],13,12); card(s,4.85,4.45,3.55,1.35,'Car Evaluation','1,728 rows · 7 headers\n6 factors + 1 decision · 4 classes',C['greenL'],C['green'],13,12); card(s,8.85,4.45,3.55,1.35,'Balance Scale','625 rows · 5 headers\n4 factors + 1 decision · 3 classes',C['orangeL'],C['orange'],13,12); takeaway(s,'老師問的「資料比數＋標頭量」現在由程式自動輸出，不再只寫 Attrs。'); footer(s,3)

s=slide(); title(s,'2. 80 / 20：最後 20% 完全不參與前面建模','同一 split 被六種方法共用；這是 final prediction set，不是 training accuracy')
steps=[('100% Dataset','原始 decision table',C['grayL'],C['gray']),('80% Train','唯一可以學習的資料',C['blueL'],C['blue']),('Reduct → Clients → LEM2 → UNION → Support / Weight','全部在 80% 內完成',C['tealL'],C['teal']),('20% Holdout','最後才拿來 prediction',C['orangeL'],C['orange']),('Evaluation','Acc · Recall · F1 · BA · Coverage',C['greenL'],C['green'])]; xs=[0.55,2.75,5.0,9.0,11.0]; widths=[1.75,1.75,3.45,1.55,1.7]
for i,(h,b,f,a) in enumerate(steps):
    rect(s,xs[i],2.0,widths[i],1.5,f); tx(s,h,xs[i]+0.08,2.28,widths[i]-0.16,0.35,12.5,a,True,PP_ALIGN.CENTER); tx(s,b,xs[i]+0.08,2.78,widths[i]-0.16,0.45,9.5,C['dark'],False,PP_ALIGN.CENTER)
    if i<len(steps)-1: line(s,xs[i]+widths[i]+0.06,2.76,xs[i+1]-0.07,2.76,C['line'],1.5)
card(s,0.9,4.15,5.55,1.6,'20% 不可以做的事','× reduct selection\n× local LEM2 induction\n× UNION support construction\n× voting weight fitting',C['redL'],C['red'],14,12); card(s,6.85,4.15,5.55,1.6,'20% 唯一做的事','✓ matching / prediction\n✓ Accuracy / Recall / Macro-F1 / Balanced Accuracy\n✓ Overall / Target / Target Candidate Coverage',C['greenL'],C['green'],14,11.7); takeaway(s,'所有 reported prediction metrics 都只看 final 20%；不存在 test leakage。'); footer(s,4)

s=slide(); title(s,'3. 同一個 frozen rule artifact 跑六種方法','前四個是既有 baseline；後兩個是本研究的 experimental variants')
items=[('FIRST_MATCH','順序：第一條 matching rule',C['orangeL'],C['orange']),('SIMPLE','rule count：一條 rule 一票',C['goldL'],C['gold']),('STANDARD','support occurrences：Σ|S(r)|',C['blueL'],C['blue']),('GLOBAL','unique class coverage：|∪S(r)|/N_d',C['greenL'],C['green']),('POPULATION STANDARD','unique population：U_d',C['tealL'],C['teal']),('POPULATION GLOBAL','(U_d/N_d) × (U_d/T_d)',C['redL'],C['red'])]
for i,(h,b,f,a) in enumerate(items):
    col=i%2; row=i//2; card(s,0.85+col*6.05,1.45+row*1.48,5.55,1.17,h,b,f,a,13.5,11.7)
tx(s,'Established baselines',1.0,5.95,4.8,0.3,12,C['gray'],True,PP_ALIGN.CENTER); tx(s,'Project-defined variants',7.1,5.95,4.8,0.3,12,C['red'],True,PP_ALIGN.CENTER); takeaway(s,'六種方法的比較必須在相同 20% test、相同 global rules、相同 conflict cases 上進行。'); footer(s,5)

policy_order=['FIRST_MATCH','SIMPLE_VOTING','STANDARD_VOTING','GLOBAL_STRENGTH','POPULATION_STANDARD','POPULATION_GLOBAL']
def metric_slide(ds, num):
    s=slide(); title(s,f'4. {ds}：六種方法的 20% Prediction Metrics','每個 cell = Accuracy / Balanced Accuracy(=Macro Recall) / Macro-F1'); part=metrics[metrics.dataset==ds].copy(); rows=[]
    for p in policy_order:
        row={'Method':short_policy(p)}
        for c in [2,3,4,5]:
            r=part[(part.policy==p)&(part.client_count==c)].iloc[0]; row[f'C={c}']=f"{r.accuracy:.3f}\n{r.balanced_accuracy:.3f}\n{r.macro_f1:.3f}"
        rows.append(row)
    df=pd.DataFrame(rows); add_table(s,df,0.72,1.45,11.9,4.15,[1.65,2,2,2,2],font_size=11.3,first_col_bold=True); tx(s,'cell 順序：Acc / BA(=Recall_macro) / Macro-F1',0.9,5.83,5.3,0.25,10,C['gray']); best=[]
    for c in [2,3,4,5]:
        q=part[part.client_count==c]; mx=q.balanced_accuracy.max(); names='/'.join(short_policy(x) for x in q[np.isclose(q.balanced_accuracy,mx)].policy); best.append(f'C{c}: {names} {mx:.3f}')
    tx(s,'BA best → '+' · '.join(best),6.1,5.78,6.05,0.35,10.5,C['teal'],True,PP_ALIGN.RIGHT); takeaway(s,'Accuracy 與 class-balanced 指標可能選出不同 winner；不能只報一個 Accuracy。'); footer(s,num)
metric_slide('Hayes-Roth',6); metric_slide('Car Evaluation',7); metric_slide('Balance Scale',8)

s=slide(); title(s,'5. Coverage：要拆成三個問題','這是本次補 code 後最重要的新區分'); card(s,0.75,1.45,3.65,2.1,'Overall Coverage','整個 20% test 中，有沒有任何 rule match？\n\n|{x : M(x)≠∅}| / |U_test|',C['blueL'],C['blue'],14,12); card(s,4.85,1.45,3.65,2.1,'Target Coverage','真實屬於 class d 的 test cases 中，有沒有任何 rule match？\n\n不要求 matching rule 的 decision=d',C['tealL'],C['teal'],14,11.7); card(s,8.95,1.45,3.65,2.1,'Target Candidate Coverage','真實屬於 class d 的 test cases 中，是否至少有一條 matching rule 真的提出 d？\n\n這才偵測「true class 根本沒進 candidate set」',C['orangeL'],C['orange'],14,11.4); tx(s,'Coverage ≠ Recall',1.0,4.35,4.8,0.4,21,C['red'],True,PP_ALIGN.CENTER); tx(s,'Coverage 問「能不能進決策」；Recall 問「最後有沒有判對」。',6.0,4.38,6.0,0.42,14,C['dark'],True,PP_ALIGN.CENTER); card(s,1.2,5.05,10.9,0.85,'為什麼要多這一層？','真實 B 可能有 rule match，但 matching rules 只有 L / R。舊 Target Coverage 會算 covered；Target Candidate Coverage 才會判斷 B 沒有被產生。',C['redL'],C['red'],12.5,10.5); takeaway(s,'Candidate Generation / Coverage failure 與 Candidate Ranking / Conflict error 必須分開。'); footer(s,9)

s=slide(); title(s,'6. 三個 UCI 的實際 Coverage','三個數字都由同一 20% holdout 自動計算；六種 policy 的 coverage 相同'); cov=metrics.groupby(['dataset','client_count'],as_index=False).first()[['dataset','client_count','overall_coverage','target_coverage_macro','target_candidate_coverage_macro']].sort_values(['dataset','client_count']); cov['Dataset']=cov.dataset.str.replace(' Evaluation','').str.replace(' Scale',''); cov2=cov[['Dataset','client_count','overall_coverage','target_coverage_macro','target_candidate_coverage_macro']].rename(columns={'client_count':'C','overall_coverage':'Overall','target_coverage_macro':'Target(any rule)','target_candidate_coverage_macro':'Target(candidate)'}); add_table(s,cov2,0.75,1.38,11.85,4.55,[2,0.7,1.35,1.6,1.65],font_size=9.7,first_col_bold=True,formats={'Overall':fmt_pct,'Target(any rule)':fmt_pct,'Target(candidate)':fmt_pct}); card(s,0.9,5.98,11.55,0.45,'最大差距','Balance C=2：Target(any rule)=93.2%，但 Target(candidate)=57.5%；代表「有規則 match」不等於「真正 target class 有 candidate rule」。',C['orangeL'],C['orange'],10.7,9.5); footer(s,10)

s=slide(); title(s,'Target Candidate Coverage：真正暴露 induction / coverage gap','特別是 Balance 的 B class，舊 coverage 幾乎看不出問題')
for idx,ds in enumerate(['Balance Scale','Car Evaluation','Hayes-Roth']):
    part=coverage[coverage.dataset==ds].copy(); part['C']=part.client_count.astype(int); part['Class']=part.target_class.astype(str); part['Any']=part.target_coverage.map(fmt_pct); part['Candidate']=part.target_candidate_coverage.map(fmt_pct); disp=part[['C','Class','Any','Candidate']]; x=0.6+idx*4.25; tx(s,ds,x,1.35,3.8,0.35,13,C['dark'],True,PP_ALIGN.CENTER); add_table(s,disp,x,1.78,3.8,3.85,[0.7,1,1,1.2],font_size=8.9)
rect(s,0.78,5.82,3.45,0.48,C['redL']); tx(s,'Balance B：C2=0%、C3=0%',0.88,5.95,3.25,0.2,9.8,C['red'],True,PP_ALIGN.CENTER); rect(s,4.95,5.82,3.45,0.48,C['orangeL']); tx(s,'Car good：C2 只有 57.1%',5.05,5.95,3.25,0.2,9.8,C['orange'],True,PP_ALIGN.CENTER); rect(s,9.15,5.82,3.45,0.48,C['blueL']); tx(s,'Hayes C2：class1/2 都只有 60%',9.25,5.95,3.25,0.2,9.8,C['blue'],True,PP_ALIGN.CENTER); footer(s,11)

s=slide(); title(s,'7. Standard：老師的「3+3」問題，現在完整算到 winner','核心不是把多的 u 當負分，而是避免同一 population 被 rule occurrences 重複放大'); card(s,0.75,1.4,3.3,2.35,'A 的兩條 rules','r1→A : {u1,u2,u3}\nr2→A : {u2,u3,u4}\n\nT_A = 3+3 = 6\nU_A = 4\nP_A = 4/6 = 2/3',C['blueL'],C['blue'],14,12.5); card(s,4.45,1.4,3.3,2.35,'若 B 的 evidence','假設 T_B = 5\n而且五個 support 都不同\n\nU_B = 5\nP_B = 1',C['grayL'],C['gray'],14,12.5); card(s,8.15,1.4,4.25,2.35,'改前 vs 改後','STANDARD：A=6 > B=5 → A\n\nPOPULATION STANDARD：\nA=U_A=4 < B=U_B=5 → B',C['greenL'],C['green'],14,13); tx(s,'Standard 看 T_d（occurrences）',1.05,4.35,4.4,0.35,18,C['blue'],True,PP_ALIGN.CENTER); tx(s,'→',5.55,4.34,0.55,0.35,20,C['gray'],True,PP_ALIGN.CENTER); tx(s,'Population Standard 看 U_d（unique objects）',6.15,4.35,6.0,0.35,18,C['teal'],True,PP_ALIGN.CENTER); takeaway(s,'新的 score 不會把「多出來的 unique u」扣分；它只是把重複出現的同一 u 從多次計價改成一次。'); footer(s,12)

def selected_transition(frame, ds, tno):
    q=frame[(frame.dataset==ds)&(frame.selected_target_no==tno)]; return q.iloc[0] if len(q) else None
s=slide(); title(s,'Standard 真實回查：同一機制可以 Correct，也可以 Harm','用 20% holdout 的真實 conflict case，不再只看 toy example'); car1=selected_transition(stdtrans,'Car Evaluation',1); car_targets=targets[targets.dataset=='Car Evaluation'].sort_values('target_no'); car_harm=car_targets[car_targets.ps_status=='HARMED'].iloc[0]; car_harm_tr=stdtrans[(stdtrans.dataset=='Car Evaluation')&(stdtrans.client_count==car_harm.client_count)&(stdtrans.row_id==car_harm.row_id)].iloc[0]
def trans_body(r):
    base=f"True = {r.true}\nSTD winner = {r.old_winner}   →   PS winner = {r.new_winner}\n\nold: T={r.old_T:.0f}, U={r.old_U:.0f}, P={r.old_P:.3f}\nnew: T={r.new_T:.0f}, U={r.new_U:.0f}, P={r.new_P:.3f}\n\n{r.status}"
    return base+(f" · A(x)={r.redundancy_asymmetry:.3f}" if not pd.isna(r.redundancy_asymmetry) else '')
card(s,0.85,1.5,5.55,3.65,'Corrected case · Car Target 1',trans_body(car1),C['greenL'],C['green'],15,13); card(s,6.9,1.5,5.55,3.65,f'Harmed case · Car Target {int(car_harm.target_no)}',trans_body(car_harm_tr),C['redL'],C['red'],15,13); card(s,1.0,5.45,11.25,0.55,'Interpretation','同樣是去重：如果錯誤 competitor 的 redundancy 高，可能救回 true class；如果 true class 的 evidence 本身高度重疊，就可能 over-correct。',C['grayL'],C['gray'],11.2,9.8); takeaway(s,'Redundancy ≠ Noise。去重本身是明確行為，但「重複 evidence 是否該被懲罰」取決於它出現在誰身上。'); footer(s,13)

s=slide(); title(s,'8. Global：原 GlobalStrength 再乘 Population Nonredundancy','不是把 numerator / denominator 真正倒過來；而是加一個 P_d=U_d/T_d correction'); card(s,0.8,1.45,3.65,2.35,'原 GlobalStrength','GS_d = U_d / N_d\n\nU_d：unique support objects\nN_d：training decision class size',C['greenL'],C['green'],14,13); card(s,4.85,1.45,3.65,2.35,'Population factor','P_d = U_d / T_d\n\nT_d：support occurrences\nP_d 越低 → evidence overlap 越高',C['orangeL'],C['orange'],14,13); card(s,8.9,1.45,3.65,2.35,'Population Global','PG_d = GS_d × P_d\n\n= U_d² / (N_d T_d)',C['redL'],C['red'],14,14); tx(s,'原本回答：「unique population 佔 class 多大？」',1.0,4.45,5.25,0.4,15,C['green'],True,PP_ALIGN.CENTER); tx(s,'＋',6.35,4.43,0.6,0.4,22,C['gray'],True,PP_ALIGN.CENTER); tx(s,'新增回答：「這些 evidence 有多少不是重複計價？」',7.0,4.45,5.3,0.4,15,C['orange'],True,PP_ALIGN.CENTER); takeaway(s,'Population Global 是 experimental correction；GlobalStrength 本身仍保留為 baseline，不被覆蓋。'); footer(s,14)

s=slide(); title(s,'Global 真實回查：winner 怎麼被 P_d 改變','Car 的 corrected / harmed case 直接算出 old → new'); car_pg=targets[(targets.dataset=='Car Evaluation')&(targets.pg_status=='CORRECTED')].iloc[0]; car_pg_tr=globtrans[(globtrans.dataset=='Car Evaluation')&(globtrans.client_count==car_pg.client_count)&(globtrans.row_id==car_pg.row_id)].iloc[0]; car_pgh=targets[(targets.dataset=='Car Evaluation')&(targets.pg_status=='HARMED')].iloc[0]; car_pgh_tr=globtrans[(globtrans.dataset=='Car Evaluation')&(globtrans.client_count==car_pgh.client_count)&(globtrans.row_id==car_pgh.row_id)].iloc[0]
def global_body(r): return f"True = {r.true}\nGS winner = {r.old_winner}   →   PG winner = {r.new_winner}\n\nold winner: T={r.old_T:.0f}, U={r.old_U:.0f}, P={r.old_P:.3f}\nnew winner: T={r.new_T:.0f}, U={r.new_U:.0f}, P={r.new_P:.3f}\n\nwinner redundancy drop = {r.winner_redundancy_drop:.3f}\n{r.status}"
card(s,0.85,1.5,5.55,3.75,f'Corrected · Car Target {int(car_pg.target_no)}',global_body(car_pg_tr),C['greenL'],C['green'],15,12.7); card(s,6.9,1.5,5.55,3.75,f'Harmed · Car Target {int(car_pgh.target_no)}',global_body(car_pgh_tr),C['redL'],C['red'],15,12.7); card(s,1.0,5.55,11.25,0.48,'Key','所有 winner change 都往 redundancy 較低的 decision 移動；但較低 redundancy 不等於一定是真實類別。',C['grayL'],C['gray'],10.8,9.4); takeaway(s,'PG 的行為符合設計；真正的研究問題是何時這個 penalty 對「錯的人」施加、何時對「對的人」施加。'); footer(s,15)

status_map={'CORRECTED':'Corrected','HARMED':'Harmed','UNCHANGED':'Unchanged','CHANGED_SAME_CORRECTNESS':'Changed / same correctness'}
def target_slide(ds,num):
    s=slide(); title(s,f'9. {ds}：Target 1–5 全部回查','不是挑幾筆漂亮案例；5 個 selected holdout targets 全部列出'); part=targets[targets.dataset==ds].sort_values('target_no').copy(); rows=[]
    for _,r in part.iterrows():
        factor_cols=[c for c in part.columns if c.startswith('factor__') and not pd.isna(r[c])]; fac=', '.join(f"{c.replace('factor__','')}={r[c]}" for c in factor_cols); fac=fac if len(fac)<=60 else fac[:57]+'…'; rows.append({'T':int(r.target_no),'C':int(r.client_count),'True':str(r.true),'Factors':fac,'STD→PS':status_map.get(r.ps_status,r.ps_status),'GS→PG':status_map.get(r.pg_status,r.pg_status),'True cand?':'YES' if bool(r.true_candidate_present) else 'NO'})
    df=pd.DataFrame(rows); add_table(s,df,0.52,1.35,12.3,4.6,[0.5,0.5,0.7,4.3,1.4,1.4,0.9],font_size=8.6,first_col_bold=True); missing=(~part.true_candidate_present.astype(bool)).sum(); corr=(part.ps_status.eq('CORRECTED').sum()+part.pg_status.eq('CORRECTED').sum()); harm=(part.ps_status.eq('HARMED').sum()+part.pg_status.eq('HARMED').sum()); tx(s,f'這 5 筆：corrected events={corr} · harmed events={harm} · true-candidate missing={missing}',0.9,6.02,11.5,0.28,11,C['teal'],True,PP_ALIGN.CENTER); footer(s,num)
target_slide('Balance Scale',16); target_slide('Car Evaluation',17); target_slide('Hayes-Roth',18)

s=slide(); title(s,'10. 908 個 Conflict Occurrences：把 error 拆開','Population correction 只改 43 個 borderline cases；另外有 45 個 conflict cases 的 true class 根本不在 candidate set'); rows=[]
for ds in ['Balance Scale','Car Evaluation','Hayes-Roth']:
    e=errors[errors.dataset==ds]; rows.append({'Dataset':ds,'Conflict':len(e),'True cand absent':int((~e.true_candidate_present).sum()),'PS corrected':int((e.ps_status=='CORRECTED').sum()),'PS harmed':int((e.ps_status=='HARMED').sum()),'PG corrected':int((e.pg_status=='CORRECTED').sum()),'PG harmed':int((e.pg_status=='HARMED').sum())})
df=pd.DataFrame(rows); add_table(s,df,0.65,1.45,12.0,2.65,[2,1,1.35,1.1,1,1.1,1],font_size=10.1,first_col_bold=True)
for i,(big,label,f,a) in enumerate([('908','conflict occurrences',C['orangeL'],C['orange']),('43','population-induced winner changes',C['redL'],C['red']),('45','conflict cases without true candidate',C['blueL'],C['blue']),('43/43','changed winner → lower redundancy',C['greenL'],C['green'])]):
    x=0.75+i*3.08; rect(s,x,4.55,2.65,1.25,f); tx(s,big,x+0.1,4.75,2.45,0.38,23,a,True,PP_ALIGN.CENTER); tx(s,label,x+0.14,5.20,2.37,0.28,9.7,C['dark'],False,PP_ALIGN.CENTER)
takeaway(s,'Coverage / induction failure 不能用 voting 修；Ranking error 才是 Standard / Global / Population policies 能處理的範圍。'); footer(s,19)

s=slide(); title(s,'11. Redundancy Asymmetry：現在已正式寫進 code 與 CSV','這是 mechanism diagnostic，不是既有文獻 metric，也不是獨立 causal proof'); tx(s,'R_d = 1 - P_d',0.9,1.55,3.2,0.5,22,C['dark'],True,PP_ALIGN.CENTER); tx(s,'A(x) = R_competitor - R_true',4.15,1.55,5.0,0.5,22,C['teal'],True,PP_ALIGN.CENTER); card(s,0.9,2.45,5.45,1.55,'A(x) > 0 · Corrected-side mechanism','錯誤 competitor 的 evidence 比 true decision 更重複。\n目前 corrected transitions 的 A(x) 全部為正。',C['greenL'],C['green'],14,12); card(s,6.85,2.45,5.45,1.55,'A(x) < 0 · Over-correction-side mechanism','true decision 本身的 evidence 比錯誤 competitor 更重複。\n目前 harmed transitions 的 A(x) 全部為負。',C['redL'],C['red'],14,12); changed=asym[asym.status.isin(['CORRECTED','HARMED'])]; statrows=[]
for method in ['POPULATION_STANDARD','POPULATION_GLOBAL']:
    for status in ['CORRECTED','HARMED']:
        q=changed[(changed.method==method)&(changed.status==status)].redundancy_asymmetry.dropna(); statrows.append({'Method':short_policy(method),'Status':status,'n':len(q),'mean A':q.mean(),'min':q.min(),'max':q.max()})
statdf=pd.DataFrame(statrows); add_table(s,statdf,1.25,4.35,10.8,1.25,[1.6,1.5,0.6,1,1,1],font_size=9.5,first_col_bold=True,formats={'mean A':lambda v:f'{v:+.3f}','min':lambda v:f'{v:+.3f}','max':lambda v:f'{v:+.3f}'}); takeaway(s,'A(x) 的 sign 在 changed cases 中是機制一致性檢查；下一步要在 repeated splits 中測它能否事前預測 corrected / harmed。'); footer(s,20)

s=slide(); title(s,'12. 老師上次問題：現在逐題的最終答案','這頁就是開會最後要交卷的版本'); answers=[('資料量？','132/5、1728/7、625/5（Rows / Headers），並明列 factors、classes、80/20。'),('20% 有沒有參與？','沒有。20% 只做最後 prediction / evaluation；所有 rules / supports / weights 只由 80% 建立。'),('六種方法結果？','72 組完整 Acc / Recall(=BA) / Macro-F1 已自動輸出；沒有 universal winner。'),('Coverage？','Overall、Target(any rule)、Target Candidate(true class represented) 三層已分開。'),('Standard 3+3？','T_d 會重複計 support occurrences；PS 改用 U_d，只讓同一 supporting object 計一次。'),('Global 怎麼改？','保留 GS=U/N，再乘 P=U/T；PG=U²/(NT)。baseline 不被覆蓋。'),('Target 1–5？','三個 dataset 各 5 筆全部回查；可分 corrected、harmed、coverage failure。')]
for i,(q,a) in enumerate(answers):
    y=1.35+i*0.69; rect(s,0.75,y,2.05,0.48,C['tealL']); tx(s,q,0.82,y+0.09,1.9,0.22,10.8,C['teal'],True,PP_ALIGN.CENTER); tx(s,a,3.05,y+0.06,9.45,0.35,11.1,C['dark'])
takeaway(s,'老師要求的每一題現在都有對應 code、CSV evidence 與簡報頁，不再靠人工口頭補。'); footer(s,21)

s=slide(); title(s,'13. Code / Evidence Closure','這次不是只改 PPT；缺的 diagnostics 已正式進 research branch'); card(s,0.75,1.45,3.75,2.05,'Code 補齊','✓ Target Candidate Coverage\n✓ Redundancy Asymmetry helper\n✓ Coverage vs Ranking error classification\n✓ Standard / Global transition decomposition',C['blueL'],C['blue'],14,11.8); card(s,4.8,1.45,3.75,2.05,'Teacher CSVs','✓ dataset metadata\n✓ six-policy metrics\n✓ coverage by class\n✓ target 1–5\n✓ STD/GS transitions\n✓ error/asymmetry tables',C['greenL'],C['green'],14,11.4); card(s,8.85,1.45,3.75,2.05,'Verification','TDD RED：缺 function / module 時 CI fail\nGREEN：implementation 後\n118 tests passed · compile PASS\nbacktrace workflow SUCCESS',C['orangeL'],C['orange'],14,11.5); card(s,0.9,4.05,11.55,1.45,'Repository state','Branch: research/conflict-client-sweep-20261020\nPR #5: Draft / Open / Not merged\nFrozen production main remains unchanged.',C['grayL'],C['gray'],13,12); takeaway(s,'這一輪完成的是 professor-question evidence layer，不修改 frozen main，也不新增第七種 prediction policy。'); footer(s,22)

s=slide(); title(s,'14. 現在可以下什麼結論、還不能下什麼結論','把「已證明的 code 行為」與「下一步要統計驗證的研究假說」分開'); card(s,0.8,1.45,5.8,3.4,'目前可以確定','• Population correction 會降低 duplicated support 對 winner 的影響。\n• 43/43 changed winners 都往 lower-redundancy decision 移動。\n• 它可以 corrected，也可以 harmed。\n• true class 不在 candidate set 時，任何 voting 都救不了。\n• Target Candidate Coverage 能直接量到這個缺口。',C['greenL'],C['green'],15,12); card(s,6.75,1.45,5.8,3.4,'目前不能直接宣稱','• Population method universally better。\n• Client 數越多就一定越好／越差。\n• Redundancy 本身就是 noise。\n• A(x) 已經是 causal predictor。\n• seed=42 的 72 rows 可直接當獨立樣本做顯著性。',C['redL'],C['red'],15,12); card(s,1.0,5.25,11.3,0.65,'下一步','Repeated matched seeds / splits → 每個 conflict case 記錄 candidate coverage、T/U/P、score margin、A(x) → Friedman / paired Wilcoxon（block-level）。',C['blueL'],C['blue'],11.2,10); takeaway(s,'下一輪不加新方法；先驗證「什麼 support structure 會讓 Population correction corrected 或 over-correct」。'); footer(s,23)

prs.save(PPTX)
print(PPTX)
print('slides=', len(prs.slides), 'bytes=', PPTX.stat().st_size)
