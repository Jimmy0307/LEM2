from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pathlib import Path

OUT = Path('docs/research/conflict_client_sweep_20261020/presentation')
OUT.mkdir(parents=True, exist_ok=True)
PPTX = OUT / 'Federated_LEM2_Final_Research_Update_20261023.pptx'

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

C = {
    'dark':'1F2933','teal':'2B7A78','tealL':'E7F2F1','blue':'4A78A8','blueL':'EAF2F8',
    'orange':'C96A3D','orangeL':'F8EAE3','green':'4C8A61','greenL':'E8F3E8',
    'red':'B65252','redL':'FBEAEA','gray':'6B7780','grayL':'F5F7F8','line':'D9E2E7','gold':'A98025','goldL':'F6F0DC'
}

def rgb(h): return RGBColor.from_string(h)

def tx(slide, text, x,y,w,h, size=14, color='1F2933', bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame; tf.clear(); tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=valign
    p=tf.paragraphs[0]; p.alignment=align
    r=p.add_run(); r.text=text; r.font.name='Aptos'; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=rgb(color)
    return box

def rect(slide,x,y,w,h,fill,line=None,round=True):
    shp=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb=rgb(fill)
    shp.line.color.rgb=rgb(line or fill)
    return shp

def line(slide,x1,y1,x2,y2,color='AAB4BA',width=1.5):
    ln=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    ln.line.color.rgb=rgb(color); ln.line.width=Pt(width)
    return ln

def title(slide,t,sub=None):
    tx(slide,t,0.62,0.38,12.0,0.55,25,'1F2933',True)
    if sub: tx(slide,sub,0.64,0.98,12.0,0.38,14.5,'2B7A78')

def footer(slide,n):
    line(slide,0.55,7.08,12.75,7.08,'D9E2E7',0.8)
    tx(slide,'Federated LEM2 · 2026-10 · research/conflict-client-sweep-20261020',0.6,7.13,10.6,0.2,8.5,'6B7780')
    tx(slide,str(n),12.3,7.13,0.35,0.2,8.5,'6B7780',False,PP_ALIGN.RIGHT)

def card(slide,x,y,w,h,head,body,fill='F5F7F8',accent='2B7A78',bs=12.5):
    rect(slide,x,y,w,h,fill)
    rect(slide,x+0.12,y+0.14,0.06,h-0.28,accent,accent,False)
    tx(slide,head,x+0.34,y+0.16,w-0.5,0.32,15, '1F2933', True)
    tx(slide,body,x+0.34,y+0.58,w-0.55,h-0.7,bs,'1F2933')

def takeaway(slide,text):
    rect(slide,0.72,6.25,11.9,0.6,'E7F2F1')
    tx(slide,'這張圖要證明：',0.94,6.42,1.35,0.2,11,'2B7A78',True)
    tx(slide,text,2.28,6.37,9.95,0.27,11.3,'1F2933')

def stat(slide,x,y,w,h,big,label,fill,accent):
    rect(slide,x,y,w,h,fill)
    tx(slide,big,x+0.1,y+0.2,w-0.2,0.48,26,accent,True,PP_ALIGN.CENTER)
    tx(slide,label,x+0.15,y+0.78,w-0.3,h-0.88,10.5,'1F2933',False,PP_ALIGN.CENTER)

def slide():
    s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb('FFFFFF'); return s

# 1
s=slide(); tx(s,'RESEARCH QUESTION',0.65,0.48,4,0.25,12,'6B7780',True)
tx(s,'Federated LEM2：Rule Conflict、Population Correction\n與 Redundancy Asymmetry',0.65,1.15,11.8,1.05,28,'1F2933',True)
tx(s,'Federated LEM2 中，client fragmentation 如何透過 rule structure、support redundancy 與 conflict，改變 prediction policy 的結果？',0.68,2.45,11.5,0.7,16.5,'2B7A78')
labels=['Client Fragmentation','Rule Structure','Conflict','Support Redundancy','Redundancy Asymmetry','Prediction Policy','Performance']
x=0.75
for i,t in enumerate(labels):
    w=1.7 if i==0 else 1.48; rect(s,x,4.35,w,0.42,['E7F2F1','EAF2F8','F8EAE3','F6F0DC','FBEAEA','E8F3E8','F5F7F8'][i]); tx(s,t,x+0.05,4.46,w-0.1,0.18,9.5,'1F2933',True,PP_ALIGN.CENTER); x+=w+0.34

tx(s,'本輪新增：6-policy experiment · 15 target backtrace · 908 conflict occurrences',0.85,5.35,11.3,0.42,15,'1F2933',True,PP_ALIGN.CENTER)
tx(s,'PR #5 維持 Draft / Open / Not merged',0.85,5.86,11.3,0.3,11,'6B7780',False,PP_ALIGN.CENTER); footer(s,1)

# 2
s=slide(); title(s,'這次更新了什麼','從「四種 policy 比較」往下挖到 population-level mechanism')
stat(s,0.75,1.65,2.55,1.55,'80 / 20','20% holdout 不參與 reduct / LEM2 / support / weight','EAF2F8','4A78A8')
stat(s,3.55,1.65,2.55,1.55,'6','4 baseline + 2 population variants','E8F3E8','4C8A61')
stat(s,6.35,1.65,2.55,1.55,'15','三個資料集各 5 個 target cases','F8EAE3','C96A3D')
stat(s,9.15,1.65,2.55,1.55,'908','C=2~5 的全部 conflict occurrences','FBEAEA','B65252')
card(s,0.9,3.65,5.6,1.9,'新增核心量','T_d = support occurrences\nU_d = unique supporting population\nP_d = U_d / T_d','F5F7F8','2B7A78',14)
card(s,6.85,3.65,5.1,1.9,'研究焦點改變','不是再加更多 voting，而是回答：\n「同一批 population 被多條 rule 重複計價時，何時修正有效？何時會 over-correct？」','F5F7F8','B65252',13.5)
takeaway(s,'研究問題已從 policy ranking 收斂到 support redundancy 與 redundancy asymmetry。'); footer(s,2)

# 3
s=slide(); title(s,'End-to-End Pipeline','六種方法共用同一個 frozen artifact，只在最後 prediction layer 分岔')
steps=[('Decision\nTable','attrs + decision'),('80/20\nSplit','20% untouched'),('Client\nPartition','C=2,3,4,5'),('Local\nLEM2','local rules'),('UNION\nAggregation','R_G=∪R_c'),('Conflict\nAudit','prediction-time'),('6 Policies','same test set')]
xs=[0.65,2.42,4.18,5.95,7.72,9.49,11.18]
for i,(h,b) in enumerate(steps):
    fill=['F5F7F8','EAF2F8','E7F2F1','F5F7F8','E7F2F1','F8EAE3','E8F3E8'][i]; rect(s,xs[i],2.1,1.45,1.35,fill)
    tx(s,h,xs[i]+0.08,2.34,1.29,0.48,13,'1F2933',True,PP_ALIGN.CENTER); tx(s,b,xs[i]+0.08,2.95,1.29,0.22,9,'1F2933',False,PP_ALIGN.CENTER)
    if i<6: line(s,xs[i]+1.48,2.77,xs[i+1]-0.06,2.77)
card(s,0.85,4.05,3.55,1.55,'Training only','reduct · discretisation · local LEM2 · support · weights','EAF2F8','4A78A8',12.2)
card(s,4.75,4.05,3.55,1.55,'Aggregation only','UNION 保留 antecedent + decision identity；不處理 conflict','E7F2F1','2B7A78',12.2)
card(s,8.65,4.05,3.55,1.55,'Holdout only','prediction · Accuracy · Recall · F1 · BA · Overall / Target Coverage','E8F3E8','4C8A61',12.0)
takeaway(s,'C 是外層實驗因素；UNION 只做 aggregation；conflict resolution 發生在 prediction policy。'); footer(s,3)

# 4
s=slide(); title(s,'Conflict：先分 Candidate Generation，再分 Candidate Ranking','有些錯誤不是 voting 問題，而是 true class 根本沒有 matching rule')
tx(s,'M(x) = { r ∈ R_G : x matches Ante(r) }',0.85,1.65,5.2,0.4,19,'1F2933',True)
tx(s,'C_conflict = { x : |{ d_r : r ∈ M(x) }| > 1 }',0.85,2.2,6.2,0.42,18,'C96A3D',True)
card(s,0.85,3.0,5.25,2.25,'Layer A · Candidate Generation / Coverage','true class 是否至少有一條 matching rule？\n\n若 true class 不在 candidate decisions 中，任何 voting / population weighting 都救不了。','EAF2F8','4A78A8',13)
card(s,6.55,3.0,5.25,2.25,'Layer B · Candidate Ranking / Conflict Resolution','多個 candidate decisions 同時存在時，哪個應該贏？\n\n六種 prediction methods 的差異只發生在這一層。','F8EAE3','C96A3D',13)
takeaway(s,'先區分「沒有正確 candidate」與「candidate 排序錯」，才能知道該修 induction 還是 scoring。'); footer(s,4)

# 5
s=slide(); title(s,'四個既有 Prediction Policies','四個 baseline 使用四種不同層次的 evidence')
items=[('FIRST_MATCH','Order · 第一條 matching rule 決定答案 · Woźniak et al. (2026)','F8EAE3','C96A3D'),('SIMPLE_VOTING','Rule Count · one matching rule = one vote · Lazo-Cortés et al. (2020)','F6F0DC','A98025'),('STANDARD_VOTING','Support Sum · Σ |S(r)| · support overlap 會重複計數','EAF2F8','4A78A8'),('GLOBAL_STRENGTH','Unique Coverage · |∪S(r)| / N_d · Bazan et al. (2000)','E8F3E8','4C8A61')]
for i,(h,b,f,a) in enumerate(items): card(s,0.85,1.55+i*1.18,11.15,0.95,h,b,f,a,11.6)
tx(s,'Order  →  Rule Count  →  Rule Support Sum  →  Unique Class Coverage Ratio',1.0,6.2,11,0.4,16,'2B7A78',True,PP_ALIGN.CENTER); footer(s,5)

# 6
s=slide(); title(s,'為什麼 Standard 會「重複放大」support','同一 training object 若被多條 same-decision rules cover，會被重複加進 score')
card(s,0.85,1.55,3.3,2.45,'Matching rules for class A','r1 → A : {u1, u2, u3}\nr2 → A : {u2, u3, u4}\n\nu2、u3 同時支持 r1 與 r2','F5F7F8','6B7780',13.5)
card(s,4.55,1.55,3.3,2.45,'STANDARD_VOTING','|S(r1)| = 3\n|S(r2)| = 3\n\nScore_STD(A|x) = 3 + 3 = 6','EAF2F8','4A78A8',14.5)
card(s,8.25,1.55,3.3,2.45,'Unique population','S(r1) ∪ S(r2) = {u1,u2,u3,u4}\n\n真正不同的 supporting objects = 4','E8F3E8','4C8A61',13.5)
tx(s,'重複量 = T_A - U_A = 6 - 4 = 2',1.0,4.55,11,0.48,22,'B65252',True,PP_ALIGN.CENTER)
tx(s,'問題不是「support 大不大」，而是「support occurrences 中有多少其實是同一批 population 重複出現」。',1.2,5.25,10.6,0.55,13.5,'1F2933',False,PP_ALIGN.CENTER)
takeaway(s,'Standard 以 rule 為加總單位，因此 support overlap 會造成 evidence inflation。'); footer(s,6)

# 7
s=slide(); title(s,'Population Correction：T / U / P','兩個方法是 project-defined experimental variants，不宣稱來自既有文獻')
card(s,0.8,1.55,3.45,1.65,'T_d · Support Occurrences','T_d = Σ |S_G(r)|\n保留每一條 rule 的 raw support 次數','EAF2F8','4A78A8',14)
card(s,4.5,1.55,3.45,1.65,'U_d · Unique Population','U_d = | ∪ S_G(r) |\n同一 training object 只算一次','E8F3E8','4C8A61',14)
card(s,8.2,1.55,3.45,1.65,'P_d · Nonredundancy','P_d = U_d / T_d\n例：4 / 6 = 2 / 3','F8EAE3','C96A3D',14)
card(s,1.2,3.75,4.9,1.6,'POPULATION_STANDARD','Score_PS(d|x) = U_d\n把 Standard 的 duplicate support inflation 直接去重。','E7F2F1','2B7A78',14)
card(s,6.75,3.75,4.9,1.6,'POPULATION_GLOBAL','Score_PG(d|x) = (U_d / N_d) × (U_d / T_d)\n= U_d² / (N_d T_d)','FBEAEA','B65252',14)
takeaway(s,'Population correction 的目的不是保證更準，而是降低同一批 population 被多條 rule 重複計價的影響。'); footer(s,7)

# 8
s=slide(); title(s,'Experimental Design','三個 UCI datasets × 四個 client counts × 六種 methods = 72 組 holdout 結果')
card(s,0.8,1.55,7.2,2.35,'Dataset metadata','Hayes-Roth: 132 rows · 4 attrs · 3 classes · 106/26\nCar Evaluation: 1,728 · 6 attrs · 4 classes · 1,382/346\nBalance Scale: 625 · 4 attrs · 3 classes · 499/126','F5F7F8','6B7780',13)
stat(s,8.45,1.55,1.65,1.45,'4','C = 2,3,4,5','E7F2F1','2B7A78'); stat(s,10.25,1.55,1.65,1.45,'6','Prediction methods','E8F3E8','4C8A61'); stat(s,8.45,3.2,3.45,1.45,'72','dataset × C × method','F8EAE3','C96A3D')
card(s,0.85,4.45,5.4,1.35,'固定比較條件','同一 dataset · 同一 seed42 split · 同一 client partition · 同一 frozen rule artifact','F5F7F8','6B7780',12.5)
card(s,6.65,4.45,5.3,1.35,'最後只看 20%','Accuracy · Recall · Macro-F1 · BA · Overall Coverage · Target Coverage','F5F7F8','6B7780',12.2)
takeaway(s,'六種方法的差異不來自資料切法，而是同一 conflict evidence 的 scoring semantics。'); footer(s,8)

# 9
s=slide(); title(s,'Six-Policy Results：沒有 Universal Winner','Population correction 在部分設定改善，但不是全面性 accuracy upgrade')
card(s,0.75,1.45,3.85,3.85,'Hayes-Roth','C2：多方法並列 0.7667\nC3：FIRST_MATCH 0.8333\nC4：FIRST_MATCH 0.7778\nC5：FIRST_MATCH 0.8333\n\nPG C5 = 0.6333，較 GS 0.7333 明顯下降','F8EAE3','C96A3D',12.3)
card(s,4.75,1.45,3.85,3.85,'Car Evaluation','C2：GLOBAL / PG = 0.7174\nC3：PG = 0.8005 最佳\nC4：GLOBAL = 0.6832\nC5：PG = 0.6338 最佳\n\nPopulation Global 在 C3、C5 有正向改善','E8F3E8','4C8A61',12.3)
card(s,8.75,1.45,3.85,3.85,'Balance Scale','C2：PS = 0.5632 最佳\nC3：STD / PS = 0.5862\nC4：STD / PS = 0.5977\nC5：GS / PG = 0.5908\n\nPopulation Standard 在部分設定改善或打平','EAF2F8','4A78A8',12.3)
takeaway(s,'Population correction 是條件式修正；效果取決於 competing decisions 的 support redundancy 結構。'); footer(s,9)

# 10
s=slide(); title(s,'Target Backtrace：從整體分數回到單筆 evidence','三個資料集各挑 5 個代表性 holdout cases，另外再掃全部 conflict cases')
for i,(n,h,b) in enumerate([('1','Factors','保留因素值'),('2','Rules','列 matching rules'),('3','Decisions','分 candidate decision'),('4','T / U / P','回查 redundancy'),('5','6 Predictions','比較 winner')]):
    x=0.75+i*2.45; sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x+0.75),Inches(1.65),Inches(0.75),Inches(0.75)); sh.fill.solid(); sh.fill.fore_color.rgb=rgb('E7F2F1' if i<3 else 'F8EAE3'); sh.line.color.rgb=rgb('2B7A78' if i<3 else 'C96A3D'); tx(s,n,x+0.75,1.85,0.75,0.22,16,'2B7A78' if i<3 else 'C96A3D',True,PP_ALIGN.CENTER); tx(s,h,x+0.15,2.62,1.95,0.3,13.5,'1F2933',True,PP_ALIGN.CENTER); tx(s,b,x+0.15,3.04,1.95,0.45,10.5,'6B7780',False,PP_ALIGN.CENTER)
stat(s,2.0,4.55,2.7,1.25,'15','selected target cases','EAF2F8','4A78A8'); stat(s,5.25,4.55,2.7,1.25,'908','all conflict occurrences','F8EAE3','C96A3D'); stat(s,8.5,4.55,2.7,1.25,'43','winner changes from PS + PG','FBEAEA','B65252')
takeaway(s,'目標不是再追一個更高 BA，而是解釋「為什麼這一筆被修正、另一筆被修壞」。'); footer(s,10)

# 11
s=slide(); title(s,'Balance Scale：同一去重機制，可以救對也可以修壞','關鍵不是 redundancy 是否存在，而是 redundancy 落在正確 decision 還是錯誤 competitor')
card(s,0.8,1.55,5.55,3.6,'Target 1 · CORRECTED','True = R\n\nStandard：L=19 > R=18 → 錯\nL：T=19, U=16\nR：T=18, U=18\n\nPS：L=16 < R=18 → 正確 R','E8F3E8','4C8A61',14)
card(s,6.75,1.55,5.55,3.6,'Target 3 · HARMED','True = L\n\nStandard：L=21 > R=19 → 正確\nL：T=21, U=17\nR：T=19, U=19\n\nPS：L=17 < R=19 → 錯誤 R','FBEAEA','B65252',14)
takeaway(s,'重複 evidence 不等於錯誤 evidence；若 true decision 本身高度重複，Population correction 可能 over-correct。'); footer(s,11)

# 12
s=slide(); title(s,'Car Evaluation：最清楚的成功與 Over-correction 例子','Population correction 會重新排序 candidate decisions，而不是增加新的 rules')
card(s,0.75,1.45,3.85,4.1,'Target 1 · PS CORRECTED','True = vgood\nStandard：good=9，vgood=9\ngood：T=9, U=7\nvgood：T=9, U=9\n\nPS：good=7 < vgood=9\n→ 修成 vgood','E8F3E8','4C8A61',12.6)
card(s,4.75,1.45,3.85,4.1,'Target 2 · PG CORRECTED','True = good\nGS(acc)=0.1140 > GS(good)=0.0909\nacc：T=50, U=35, P=0.70\nPG(acc)=0.0798\nPG(good)=0.0909\n→ acc 改判 good','EAF2F8','4A78A8',12.6)
card(s,8.75,1.45,3.85,4.1,'Target 4 · PG HARMED','True = acc\nGS(acc)=0.1270 > GS(unacc)=0.1095\nacc：T=67, U=39, P=0.5821\nPG(acc)=0.0739\n→ 被 unacc 反超','FBEAEA','B65252',12.6)
takeaway(s,'如果 repeated support 本來就是有效證據，Population penalty 就可能造成 over-correction。'); footer(s,12)

# 13
s=slide(); title(s,'Hayes-Roth：小樣本下的 Over-correction 風險最明顯','重要反例：不能把 Population 方法包裝成 universal improvement')
card(s,0.85,1.6,5.4,2.3,'Target 1 · PS 成功','True = class 3\n錯誤 class 1：T=10 → U=7\n真實 class 3：T=9 → U=9\n\nPS：7 < 9 → 修正成功','E8F3E8','4C8A61',14)
card(s,6.65,1.6,5.4,2.3,'Target 2 · PS 修壞','True = class 1\n真實 class 1：T=15 → U=8\nclass 2：T=9 → U=9\n\nPS：8 < 9 → 原正確結果被翻掉','FBEAEA','B65252',14)
stat(s,1.6,4.45,3.2,1.15,'-10.0 pp','Hayes C5：GS 0.7333 → PG 0.6333','FBEAEA','B65252')
card(s,5.25,4.35,6.5,1.35,'目前只能形成機制假說','small local sample + high overlap 可能讓 redundancy penalty 過度懲罰有效 support；需要 repeated seeds / splits 才能正式驗證。','F8EAE3','C96A3D',12.4)
takeaway(s,'方法行為合理，不代表方法在每個資料條件下都會提升 performance。'); footer(s,13)

# 14
s=slide(); title(s,'有些 Case 根本不是 Scoring 問題','若 true class 沒有 matching rule，六種方法都只是在錯誤 candidates 中重新排序')
card(s,0.85,1.55,5.35,2.4,'Balance Target 5','True = B\nCandidate decisions = {L, R}\n\nB 不在 matching-rule decisions 中。\nPG 即使把 R 改成 L，兩者仍都不是 B。','EAF2F8','4A78A8',14)
card(s,6.6,1.55,5.35,2.4,'Car Target 5','True = good\nCandidate decisions = {acc, vgood}\n\ngood 沒有 matching rule。\nPopulation correction 只能換錯誤 winner。','F8EAE3','C96A3D',14)
tx(s,'Coverage / Induction Error',1.4,4.55,4.2,0.45,21,'4A78A8',True,PP_ALIGN.CENTER); tx(s,'≠',6.1,4.52,0.5,0.45,24,'6B7780',True,PP_ALIGN.CENTER); tx(s,'Conflict Ranking Error',7.0,4.55,4.2,0.45,21,'C96A3D',True,PP_ALIGN.CENTER)
tx(s,'Target Coverage 看「有沒有 rule 能處理該 class」；Recall 才看「最後有沒有判對」。',1.2,5.35,10.8,0.45,13.5,'1F2933',False,PP_ALIGN.CENTER)
takeaway(s,'先把 candidate generation 與 candidate ranking 拆開，才能避免把所有錯誤都歸咎於 voting。'); footer(s,14)

# 15
s=slide(); title(s,'908 個 Conflict Occurrences：Population 只改少數 Borderline Cases','大部分 conflict cases 的 winner 沒變；真正變動集中在 redundancy 足以翻轉排序的 cases')
card(s,0.75,1.45,8.1,4.0,'Status counts','Balance · PS：Corrected 3 / Harmed 1 / Other 1 / Unchanged 261\nBalance · PG：Corrected 3 / Harmed 3 / Other 3 / Unchanged 257\nCar · PS：Corrected 2 / Harmed 4 / Unchanged 592\nCar · PG：Corrected 7 / Harmed 6 / Other 1 / Unchanged 584\nHayes · PS：Corrected 1 / Harmed 4 / Unchanged 39\nHayes · PG：Corrected 0 / Harmed 4 / Unchanged 40','F5F7F8','6B7780',12.1)
stat(s,9.35,1.6,2.4,1.35,'908','conflict occurrences','F8EAE3','C96A3D'); stat(s,9.35,3.15,2.4,1.35,'43','winner changes','FBEAEA','B65252'); tx(s,'PS changed = 16\nPG changed = 27',9.45,4.8,2.2,0.72,13,'1F2933',True,PP_ALIGN.CENTER)
takeaway(s,'Population correction 不會重洗全部 conflict；它只影響少數 scoring margin 接近的 cases。'); footer(s,15)

# 16
s=slide(); title(s,'最重要的 Mechanism Check：43 / 43','所有被 Population correction 改變的 winner，都往 support redundancy 較低的 decision 移動')
stat(s,1.05,1.55,4.0,2.0,'43 / 43','changed winners → lower-redundancy decision','E8F3E8','4C8A61')
card(s,5.45,1.55,6.25,2.0,'這證明了什麼？','Population correction 確實降低重複 evidence 的 scoring influence。\n\n但它只證明「在做去重」，不代表「去重一定提升真實分類」。','F5F7F8','6B7780',14)
card(s,0.85,4.1,3.35,1.3,'成功修正','錯誤 competitor redundancy 高於 true decision','E8F3E8','4C8A61',11.8)
card(s,5.0,4.1,3.35,1.3,'Over-correction','true decision 自己的有效 evidence 更重複','FBEAEA','B65252',11.8)
card(s,9.15,4.1,3.35,1.3,'Coverage failure','true decision 不在 candidate set；weighting 救不了','EAF2F8','4A78A8',11.8)
takeaway(s,'方法確實執行去重，但「去重是否值得」取決於 redundancy 的位置。'); footer(s,16)

# 17
s=slide(); title(s,'真正值得研究的是 Redundancy Asymmetry','不是單看 P_d，而是比較 true decision 與 competing decision 誰更依賴重複 support')
tx(s,'A(x) = Redundancy_competitor − Redundancy_true',1.0,1.55,11,0.5,23,'1F2933',True,PP_ALIGN.CENTER)
card(s,0.85,2.55,5.25,2.15,'A(x) > 0','錯誤 competitor 的 evidence 更重複\n\nPopulation penalty 比較可能壓低錯誤 winner\n→ CORRECTED','E8F3E8','4C8A61',14)
card(s,6.55,2.55,5.25,2.15,'A(x) < 0','true decision 的 evidence 更重複\n\nPopulation penalty 可能把有效 evidence 一起壓低\n→ HARMED / over-correction','FBEAEA','B65252',14)
tx(s,'Redundancy ≠ Noise',2.6,5.15,3.1,0.45,18,'C96A3D',True,PP_ALIGN.CENTER); tx(s,'Redundancy Asymmetry = 下一個 mechanism variable',6.2,5.15,4.9,0.45,18,'2B7A78',True,PP_ALIGN.CENTER)
takeaway(s,'下一輪應用 repeated seeds 驗證 asymmetry 是否穩定預測 corrected / harmed，而不是再加第七種 voting。'); footer(s,17)

# 18
s=slide(); title(s,'目前整體結論','從「哪個 policy 最好」進一步變成「什麼 evidence structure 讓哪個 policy 有效」')
con=[('1','UNION 不解決 conflict','真正影響分類的是 prediction-time competing decisions。'),('2','沒有 universal winner','最佳 method 依 dataset、C、rule structure 與 conflict workload 改變。'),('3','Standard 的確會重複計價','support occurrences 與 unique supporting objects 是兩個不同概念。'),('4','Population 是條件式修正','能壓掉錯誤 redundancy，也可能懲罰有效 repeated evidence。'),('5','Coverage 與 Ranking 必須分開','true class 缺席時，任何 scoring policy 都無法救回。'),('6','核心是 redundancy asymmetry','43/43 changed winners 都往 redundancy 較低的 decision 移動。')]
for i,(n,h,b) in enumerate(con):
    col=0 if i<3 else 1; row=i%3; x=0.75+col*6.15; y=1.4+row*1.55
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(0.55),Inches(0.55)); sh.fill.solid(); sh.fill.fore_color.rgb=rgb('E7F2F1' if col==0 else 'F8EAE3'); sh.line.color.rgb=rgb('2B7A78' if col==0 else 'C96A3D'); tx(s,n,x,y+0.15,0.55,0.2,13,'2B7A78' if col==0 else 'C96A3D',True,PP_ALIGN.CENTER); tx(s,h,x+0.75,y-0.02,5.0,0.35,14,'1F2933',True); tx(s,b,x+0.75,y+0.45,5.0,0.58,11,'6B7780')
takeaway(s,'目前最站得住腳的貢獻不是「新方法一定更準」，而是揭示 support redundancy 的作用與限制。'); footer(s,18)

# 19
s=slide(); title(s,'下一步：把 Mechanism Evidence 變成 Statistical Evidence','目前 seed42 / single 80-20 split 可形成假說，但還不能做強因果或顯著性主張')
card(s,0.75,1.45,3.7,2.0,'1 · Repeated seeds / splits','每個 dataset × C × policy 建立分佈，而不是單一數字。','EAF2F8','4A78A8',14)
card(s,4.8,1.45,3.7,2.0,'2 · Asymmetry features','對全部 conflict cases 輸出 true / competitor redundancy、margin 與 coverage state。','E7F2F1','2B7A78',14)
card(s,8.85,1.45,3.7,2.0,'3 · Statistical tests','Friedman + post-hoc；paired Wilcoxon；避免把 test rows 當 algorithm-level replicates。','E8F3E8','4C8A61',13.5)
card(s,1.2,4.05,5.0,1.45,'Research question','Client fragmentation 如何透過 rule structure 與 redundancy asymmetry，改變 conflict-resolution policy 的 performance？','F8EAE3','C96A3D',13.2)
card(s,6.75,4.05,5.0,1.45,'Repository state','Jimmy0307/LEM2\nbranch: research/conflict-client-sweep-20261020\nPR #5: Draft · Open · Not merged','F5F7F8','6B7780',12.5)
takeaway(s,'下一步先驗證 redundancy asymmetry；不要再增加 method，避免研究主軸被 voting variants 稀釋。'); footer(s,19)

prs.save(PPTX)
print(PPTX)
