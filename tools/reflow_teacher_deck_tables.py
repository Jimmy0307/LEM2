from __future__ import annotations
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

PPT = Path('docs/research/conflict_client_sweep_20261020/presentation/Federated_LEM2_Teacher_Questions_Complete_20261023.pptx')
prs = Presentation(PPT)
FONT='Noto Sans CJK TC'
NAVY=RGBColor(39,58,76); TEXT=RGBColor(54,69,79); MUTED=RGBColor(102,117,128); WHITE=RGBColor(255,255,255)
TEAL_D=RGBColor(54,117,108); TEAL_L=RGBColor(232,243,241)
BLUE_L=RGBColor(233,241,247); BLUE=RGBColor(88,130,164); PINK_L=RGBColor(250,235,233); PINK=RGBColor(191,103,94)
GOLD_L=RGBColor(248,243,222); GOLD=RGBColor(147,119,55); GRAY_L=RGBColor(244,247,248); GRAY=RGBColor(219,226,230)
GREEN_L=RGBColor(234,245,235); GREEN=RGBColor(86,139,91); RED_L=RGBColor(252,235,233); RED=RGBColor(190,93,84)

def I(v): return Inches(v)
def remove(shape):
    el=shape._element; el.getparent().remove(el)
def rect(s,x,y,w,h,fill=WHITE,line=GRAY,round=True):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE,I(x),I(y),I(w),I(h)); sh.fill.solid(); sh.fill.fore_color.rgb=fill; sh.line.color.rgb=line; sh.line.width=Pt(.8); return sh
def text(s,t,x,y,w,h,size=10,color=TEXT,bold=False,align=PP_ALIGN.LEFT):
    b=s.shapes.add_textbox(I(x),I(y),I(w),I(h)); tf=b.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=tf.margin_right=I(.04); tf.margin_top=tf.margin_bottom=I(.03); tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=str(t); r.font.name=FONT; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; return b
def cell(s,t,x,y,w,h,fill=WHITE,color=TEXT,bold=False,size=10,align=PP_ALIGN.CENTER):
    rect(s,x,y,w,h,fill,GRAY,False); text(s,t,x+.02,y+.01,w-.04,h-.02,size,color,bold,align)
def table_matrix(shape): return [[c.text for c in row.cells] for row in shape.table.rows]
def only_tables(s): return [sh for sh in s.shapes if sh.has_table]

def slide3(s):
    tbl=only_tables(s)[0]; m=table_matrix(tbl)
    shapes=list(s.shapes)
    for idx in range(14,1,-1): remove(shapes[idx])
    by={r[0]:r for r in m[1:]}; order=[('Hayes-Roth',BLUE_L,BLUE),('Car Evaluation',TEAL_L,TEAL_D),('Balance Scale',GOLD_L,GOLD)]; xs=[.78,4.52,8.26]
    for x,(name,fill,accent) in zip(xs,order):
        r=by[name]; rect(s,x,1.55,3.45,3.95,WHITE,GRAY,True); rect(s,x+.12,1.72,.07,3.55,accent,accent,False)
        text(s,name,x+.35,1.72,2.85,.42,15,accent,True); text(s,f'{r[1]} rows',x+.35,2.25,2.85,.48,21,NAVY,True)
        yy=2.95
        for lab,val in [('Headers',r[2]),('Factors',r[3]),('Classes',r[5]),('Train 80%',r[6]),('Test 20%',r[7])]:
            text(s,lab,x+.35,yy,1.35,.33,10.2,MUTED); text(s,val,x+1.78,yy,1.05,.33,11.2,TEXT,True,PP_ALIGN.RIGHT); yy+=.46
        text(s,'Headers = factors + 1 decision',x+.35,5.02,2.85,.28,9.3,MUTED)

def metric(s):
    tbl=only_tables(s)[0]; m=table_matrix(tbl); remove(tbl)
    x0,y0=.72,1.42; widths=[1.55,2.59,2.59,2.59,2.59]; hh=.52; hr=.56; x=x0
    for j,h in enumerate(m[0]): cell(s,h,x,y0,widths[j],hh,TEAL_L,TEAL_D,True,11.5); x+=widths[j]
    for i,row in enumerate(m[1:]):
        y=y0+hh+i*hr; method=row[0]; fill=PINK_L if method in ('PS','PG') else (WHITE if i%2==0 else GRAY_L); x=x0
        cell(s,method,x,y,widths[0],hr,fill,PINK if method in ('PS','PG') else TEXT,True,11,PP_ALIGN.LEFT); x+=widths[0]
        for j,v in enumerate(row[1:]):
            vals=[q for q in v.splitlines() if q.strip()]; t='  ·  '.join([f'{k} {vals[n]}' for n,k in enumerate(['Acc','BA','F1']) if n<len(vals)])
            cell(s,t,x,y,widths[j+1],hr,fill,TEXT,False,9.4,PP_ALIGN.CENTER); x+=widths[j+1]
    rect(s,.9,5.18,11.55,.42,BLUE_L,RGBColor(207,221,231),True); text(s,'每個 cell 固定為單行：Accuracy · Balanced Accuracy (= Macro Recall) · Macro-F1',1.08,5.23,11.2,.28,10.5,BLUE,True)

def slide10(s):
    tbl=only_tables(s)[0]; m=table_matrix(tbl); remove(tbl); rows={}
    for r in m[1:]: rows.setdefault(r[0],[]).append(r[1:])
    order=[('Balance','Balance Scale',BLUE_L,BLUE),('Car','Car Evaluation',TEAL_L,TEAL_D),('Hayes-Roth','Hayes-Roth',GOLD_L,GOLD)]; xs=[.67,4.58,8.49]
    for x,(key,title,fill,accent) in zip(xs,order):
        rect(s,x,1.40,3.58,4.25,WHITE,GRAY,True); text(s,title,x+.18,1.55,3.20,.38,14,accent,True,PP_ALIGN.CENTER)
        widths=[.52,.88,.98,1.00]; xx=x+.10
        for j,h in enumerate(['C','Overall','Any rule','Candidate']): cell(s,h,xx,2.05,widths[j],.48,fill,accent,True,9.2); xx+=widths[j]
        yy=2.53
        for k,r in enumerate(rows[key]):
            xx=x+.10; rf=WHITE if k%2==0 else GRAY_L
            for j,val in enumerate(r): cell(s,val,xx,yy,widths[j],.63,rf,accent if j==3 else TEXT,j in (0,3),10); xx+=widths[j]
            yy+=.63
        text(s,'Target(candidate) = true class 是否真的進入 matching candidate set',x+.22,5.16,3.14,.33,8.7,MUTED,False,PP_ALIGN.CENTER)

def slide11(s):
    tabs=only_tables(s); mats=[table_matrix(t) for t in tabs]
    for t in tabs: remove(t)
    fills=[BLUE_L,TEAL_L,GOLD_L]; accents=[BLUE,TEAL_D,GOLD]; xs=[.60,4.85,9.10]
    for x,fill,accent,m in zip(xs,fills,accents,mats):
        grouped={}
        for r in m[1:]: grouped.setdefault(r[0],[]).append((r[1],r[3]))
        rect(s,x,1.76,3.8,3.88,WHITE,GRAY,True); yy=1.95
        for idx,c in enumerate(sorted(grouped, key=lambda z:int(float(z)))):
            pairs=grouped[c]; details=(' · '.join([f'{a} {b}' for a,b in pairs[:2]])+'\n'+' · '.join([f'{a} {b}' for a,b in pairs[2:]])) if len(pairs)==4 else ' · '.join([f'{a.replace(".0","")} {b}' for a,b in pairs])
            rect(s,x+.12,yy,3.56,.78,fill if idx%2==0 else WHITE,RGBColor(226,232,235),True); text(s,f'C{int(float(c))}',x+.24,yy+.08,.52,.62,11,accent,True,PP_ALIGN.CENTER); text(s,details,x+.78,yy+.06,2.75,.66,8.7,TEXT,False,PP_ALIGN.LEFT); yy+=.86

def target_slide(s):
    tbl=only_tables(s)[0]; m=table_matrix(tbl); remove(tbl); y0=1.36
    for idx,row in enumerate(m[1:]):
        t,c,true,factors,std,gs,cand=row; y=y0+idx*.92; base=WHITE if idx%2==0 else GRAY_L
        rect(s,.58,y,12.18,.80,base,GRAY,True); text(s,f'T{t}',.72,y+.08,.45,.27,11,NAVY,True,PP_ALIGN.CENTER); text(s,f'C={c}',.72,y+.39,.45,.23,8.8,MUTED,False,PP_ALIGN.CENTER)
        text(s,'True',1.22,y+.08,.45,.22,8.6,MUTED,False,PP_ALIGN.CENTER); text(s,true.replace('.0',''),1.22,y+.31,.45,.30,11,TEXT,True,PP_ALIGN.CENTER)
        text(s,factors.replace(', ',' · ').replace('…',''),1.78,y+.09,5.45,.60,9,TEXT,False,PP_ALIGN.LEFT)
        for (lab,val),xx,ww in zip([('STD→PS',std),('GS→PG',gs),('True cand',cand)],[7.45,9.15,11.05],[1.48,1.68,1.48]):
            if 'Corrected' in val: fill,fg=GREEN_L,GREEN
            elif 'Harmed' in val: fill,fg=RED_L,RED
            elif val=='NO': fill,fg=PINK_L,PINK
            elif 'Changed' in val: fill,fg=GOLD_L,GOLD
            else: fill,fg=BLUE_L,BLUE
            rect(s,xx,y+.10,ww,.58,fill,fill,True); text(s,lab,xx+.05,y+.13,ww-.10,.17,7.4,MUTED,False,PP_ALIGN.CENTER); text(s,val.replace('Changed / same correctness','Changed / same'),xx+.05,y+.31,ww-.10,.27,8.4,fg,True,PP_ALIGN.CENTER)

def slide19(s):
    tbl=only_tables(s)[0]; m=table_matrix(tbl); remove(tbl); colors=[(BLUE_L,BLUE),(TEAL_L,TEAL_D),(GOLD_L,GOLD)]; xs=[.68,4.56,8.44]
    for x,row,(fill,accent) in zip(xs,m[1:],colors):
        name,conf,missing,psc,psh,pgc,pgh=row; rect(s,x,1.48,3.58,2.42,WHITE,GRAY,True); text(s,name,x+.15,1.62,3.28,.35,14,accent,True,PP_ALIGN.CENTER); yy=2.08
        for lab,val in [('Conflict',conf),('True candidate absent',missing),('PS corrected / harmed',f'{psc} / {psh}'),('PG corrected / harmed',f'{pgc} / {pgh}')]:
            text(s,lab,x+.22,yy,2.15,.28,9,MUTED); text(s,val,x+2.35,yy,.95,.28,11,TEXT,True,PP_ALIGN.RIGHT); yy+=.43

def slide20(s):
    tbl=only_tables(s)[0]; m=table_matrix(tbl); remove(tbl); xs=[.96,4.02,7.08,10.14]
    for x,row in zip(xs,m[1:]):
        method,status,n,mean,mn,mx=row; fill,accent=(GREEN_L,GREEN) if 'CORRECTED' in status else (RED_L,RED)
        rect(s,x,4.38,2.36,1.18,fill,fill,True); text(s,f'{method} · {status}',x+.08,4.48,2.20,.24,9.6,accent,True,PP_ALIGN.CENTER); text(s,f'n={n}',x+.08,4.75,2.20,.22,9.2,TEXT,True,PP_ALIGN.CENTER); text(s,f'mean {mean}',x+.08,4.98,2.20,.20,8.7,TEXT,False,PP_ALIGN.CENTER); text(s,f'min {mn}  ·  max {mx}',x+.08,5.19,2.20,.20,8.2,MUTED,False,PP_ALIGN.CENTER)

slide3(prs.slides[2])
for n in (6,7,8): metric(prs.slides[n-1])
slide10(prs.slides[9]); slide11(prs.slides[10])
for n in (16,17,18): target_slide(prs.slides[n-1])
slide19(prs.slides[18]); slide20(prs.slides[19])
prs.save(PPT)
print('TABLE_REFLOW=PASS', PPT)
